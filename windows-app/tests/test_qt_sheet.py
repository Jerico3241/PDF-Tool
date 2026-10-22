"""Tabellen bearbeiten (seit 5.0 Beta 2): Rechenkern in LibreOffice, Tabellen-Ansicht (Entwurf 2-tabelle.png) und
Startbildschirm (Entwurf 3-startbildschirm.png).

Die Tests mit dem Rechenkern brauchen LibreOffice (``soffice`` im PATH oder ``NOVADESK_LIBREOFFICE``); ohne laufen nur
die übrigen. Im Setup prüft ``build.py probe_calc`` den Rechenkern im mitgelieferten LibreOffice."""

from __future__ import annotations

import datetime
from pathlib import Path

import pytest
from PySide6.QtCore import QObject, QUrl
from PySide6.QtQml import QQmlComponent

from conftest import wait_until
from qtutil import pump


def _calc_available() -> bool:
    from tools.sheet import engine

    return engine.available()


needs_calc = pytest.mark.skipif(not _calc_available(), reason="LibreOffice fehlt – Rechenkern nicht prüfbar")


def _workbook(path: Path) -> Path:
    import openpyxl
    from openpyxl.styles import Font

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Verträge"
    sheet.append(["Kunde", "Vertrag", "Beginn", "Betrag", "Status"])
    for row in (("Müller", "Wartung", datetime.date(2026, 1, 1), 480, "aktiv"), ("Schmidt", "Heizung", datetime.date(2026, 3, 15), 1250, "aktiv"),
                ("Weber", "Prüfung", datetime.date(2026, 6, 1), 190, "läuft aus"), ("Fischer", "Wartung", datetime.date(2026, 9, 1), 480, "aktiv")):
        sheet.append(list(row))
    sheet["C6"], sheet["D6"] = "Summe", "=SUM(D2:D5)"
    for cell in ("A1", "B1", "C1", "D1", "E1"):
        sheet[cell].font = Font(name="Calibri", bold=True)
    for row in range(2, 7):
        sheet[f"C{row}"].number_format = "DD.MM.YYYY"
        sheet[f"D{row}"].number_format = '#,##0.00 "€"'
    book.create_sheet("Kunden")
    book.save(path)
    return path


def _open(h, path: Path):
    reader = h.runtime.singletons["Reader"]
    reader.open_paths([str(path)])
    assert wait_until(lambda: reader.current is not None and getattr(reader.current, "ready", False), 120)
    pump(0.5)
    return reader.current


def test_splash_follows_the_design(qt_application) -> None:
    """Startbildschirm: Symbol, Name, Ladebalken (wächst nur) und »Wird gestartet …«; dunkel wie die App; geht ohne
    Hauptfenster gleich zu. Gezeichnet ohne QML – er berührt die QML-Engine nicht."""
    from PySide6.QtGui import QColor

    from qtapp.splash import DARK, Splash

    splash = Splash.show({"theme": "dark", "sprache": "de"})
    assert splash is not None
    window = splash.window
    assert window.isVisible() and window.dark and window.status == "Wird gestartet …"
    window.setGeometry(0, 0, 1280, 800)
    splash.step(0.5)
    splash.step(0.3)
    assert splash.progress == pytest.approx(0.5)
    image = window.grab().toImage()
    assert image.pixelColor(10, 10) == QColor(DARK["background"])
    bar_y = round(800 / 2 - 157 + 200 + 15 + 61 + 2)  # Mitte des Ladebalkens (wie in paintEvent)
    assert image.pixelColor(640 - 100, bar_y).blue() > 200  # gefüllt: blauer Verlauf
    assert image.pixelColor(640 + 100, bar_y) == QColor(DARK["track"])  # Rest: Spur
    splash.finish(None)
    pump(0.2)
    assert splash._closed  # noqa: SLF001


def test_sheet_view_loads_without_a_sheet(ui_app) -> None:
    """Die Tabellen-Ansicht lädt ohne Meldungen, auch ohne Tabelle (im Setup aus der Ressource wie alle Seiten)."""
    from qtapp.application import QML_DIR

    h = ui_app
    component = QQmlComponent(h.engine, QUrl.fromLocalFile(str(QML_DIR / "PdfTool" / "Reader" / "SheetView.qml")))
    assert component.isReady(), component.errorString()
    view = component.create()
    assert view is not None
    view.setParentItem(h.window.contentItem())
    pump(0.3)
    for name in ("ribbon_start", "sheetFormulaBar", "sheetGrid", "sheetTabs", "sheetOverview", "sheetPdf"):
        assert view.findChild(QObject, name) is not None, name
    view.deleteLater()
    pump(0.2)


@needs_calc
def test_edit_save_and_reopen(ui_app, tmp_path: Path) -> None:
    """Excel-Datei öffnet sich zum Bearbeiten: Eingabe, Formel mit deutschem Namen, Statuszeile mit Summe und Anzahl,
    Speichern (auch mit Diagramm danach »gespeichert«) und unverändert wieder lesbar."""
    import openpyxl

    h = ui_app
    h.app.dialogs.shutdown()
    sheet = _open(h, _workbook(tmp_path / "Verträge 2026.xlsx"))
    assert sheet.kind == "tabelle" and [entry["name"] for entry in sheet.sheets] == ["Verträge", "Kunden"]
    view = h.item("sheetView")
    assert view is not None and view.property("sheet") is sheet
    assert not h.item("readerDocumentArea").isVisible()  # die PDF-Ansicht bleibt verborgen
    # Betrag: eigenes Excel-Format mit € gilt als Währung
    sheet.select(1, 3, False)
    assert wait_until(lambda: sheet.formatKind == "currency", 10)
    # Statuszeile: Summe und Anzahl der Auswahl, rechts der Zoom
    sheet.selectRange(1, 3, 4, 3)
    assert wait_until(lambda: sheet.stats.get("count") == 4, 10)
    pump(0.3)
    texts = [item.property("text") for item in h.items("sheetZoomText")]
    assert texts == ["100 %"]
    assert sheet.stats["sum"] == "2.400,00 €"
    # Eingabe und Formel (deutsche Funktionsnamen wie in Excel)
    sheet.select(7, 0, False)
    sheet.commit('=SVERWEIS("Weber";A2:D5;4;0)', 1, 0)
    sheet.select(7, 0, False)
    assert wait_until(lambda: sheet.cellInput == '=SVERWEIS("Weber";A2:D5;4;0)', 10)
    sheet.select(8, 0, False)
    sheet.commit("Neu", 1, 0)
    assert wait_until(lambda: sheet.dirty, 10)
    sheet.selectRange(0, 3, 4, 3)
    sheet.addChart("column")
    assert wait_until(lambda: len(sheet.objects) == 1, 20)
    assert sheet.save_and_wait() and wait_until(lambda: not sheet.dirty, 5)
    book = openpyxl.load_workbook(sheet.path, data_only=True)
    assert book.sheetnames == ["Verträge", "Kunden"]
    assert book["Verträge"]["A8"].value == 190 and book["Verträge"]["A9"].value == "Neu"
    assert len(openpyxl.load_workbook(sheet.path)["Verträge"]._charts) == 1  # noqa: SLF001


@needs_calc
def test_new_sheet_with_ctrl_n(ui_app) -> None:
    """Strg+N (Kürzel »app.new«) legt eine leere Tabelle an – in der Tabellen-Ansicht mit dem Reiter »Start«."""
    h = ui_app
    h.app.dialogs.shutdown()
    reader = h.runtime.singletons["Reader"]
    shortcut = h.window.findChild(QObject, "newSheetShortcut")
    assert shortcut is not None
    reader.newSheet()
    assert wait_until(lambda: reader.current is not None and getattr(reader.current, "ready", False), 120)
    pump(0.5)
    assert reader.current.kind == "tabelle" and h.item("sheetView").property("ribbon") == "start"
