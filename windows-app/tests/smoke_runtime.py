"""Prüft eine eingebettete bzw. installierte Python-Laufzeit der App – ohne zusätzliche Pakete.

Aufruf mit der zu prüfenden Laufzeit (nicht mit einem System-Python):

    build\\payload\\runtime\\python.exe -s windows-app\\tests\\smoke_runtime.py --app build\\payload\\app
    %LOCALAPPDATA%\\PDF-Tool\\runtime\\python.exe -s smoke_runtime.py --app ...\\app --ui

Geprüft wird:
1. Import aller Laufzeitmodule (PySide6 mit Qt Quick, numpy, pandas, openpyxl, xlrd, reportlab,
   PIL, pikepdf mit qpdf, pypdfium2 mit PDFium, pypdf) und der App-Module beider Werkzeuge;
   tkinter ist nicht mehr Teil der Laufzeit (seit 2.7.0)
2. Vertragsübersichten: eine echte PDF aus einer Excel mit fett formatierter Zelle und
   formatierter Fußzeile; Kundenakte 2.0 (Übernahme einer Kundenhistorie aus 2.3 mit
   Sicherung, Wiedererkennung per E-Mail, Speichern) und Live-Vorschau (PDF im
   Hintergrund erzeugen, Seite mit PDFium zeichnen, temporäre Dateien löschen);
   Stapel: zwei Excel-Listen prüfen, Kunden erkennen bzw. ergänzen und mit derselben
   Engine nacheinander erstellen – eine vorhandene PDF wird nicht überschrieben;
   Vertragsvergleich: Stand nach dem Export speichern (unverändert nicht doppelt),
   Änderungen erkennen
3. PDF reparieren: dieselbe PDF mit beschädigter Querverweistabelle im eigenen
   Arbeitsprozess analysieren und reparieren (wie in der App), Ausgabe prüfen – auch
   ohne »_repariert« (nummeriert, das Original bleibt);
   erweiterte Wiederherstellung: klassische PDF ohne xref, Trailer, %%EOF und mit
   defektem Seitenbaum rekonstruieren, Text und Seiten mit pypdf prüfen
4. Oberfläche: die QML-Oberfläche aus der eingebauten Ressource (qml_rc) laden, ohne dass die
   QML-Engine etwas meldet; mit ``--ui`` zusätzlich der Programmstart wie per Verknüpfung
   (Hauptfenster mit Startseite, Werkzeuge und die Ansichten »Stapel« und »Vorschau« öffnen;
   Kundenakte standardmäßig aus, »Kunden« erst nach dem Einschalten – ohne Neustart;
   Einstellungen werden gespeichert). ``--offscreen`` prüft ohne sichtbares Fenster.

``--part runtime`` prüft nur 1–3 (Laufzeit), ``--part ui`` nur 1 und 4 (QML und Programmstart) –
so zeigt die CI beide Prüfungen als eigene Schritte, ohne etwas doppelt zu tun.

Endet mit Code 0 und »OK«, sonst mit einer Fehlermeldung und Code 1.
"""

from __future__ import annotations

import argparse
import os
import sys
import tempfile
import time
import traceback
from pathlib import Path


def check(condition: bool, message: str) -> None:
    if not condition:
        raise SystemExit(f"FEHLER: {message}")


def main() -> int:
    # Windows-Konsole (cp1252): nicht darstellbare Zeichen ersetzen statt abzubrechen
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0] if __doc__ else None)
    parser.add_argument("--app", required=True, help="app-Ordner der zu prüfenden Installation bzw. des Pakets")
    parser.add_argument("--ui", action="store_true", help="zusätzlich das Hauptfenster starten und schließen")
    parser.add_argument("--offscreen", action="store_true", help="Qt ohne sichtbares Fenster (Plattform »offscreen«)")
    parser.add_argument(
        "--part",
        choices=("all", "runtime", "ui"),
        default="all",
        help="runtime: Module, PDF, Kundenakte, Vorschau, Stapel, Vergleich, Reparatur (ohne Fenster); ui: Oberfläche aus der Ressource und Programmstart; all: beides",
    )
    args = parser.parse_args()

    app_dir = Path(args.app).resolve()
    check((app_dir / "start.py").is_file() and (app_dir / "qtapp" / "application.py").is_file(), f"App-Ordner nicht gefunden: {app_dir}")
    if args.offscreen:
        os.environ["QT_QPA_PLATFORM"] = "offscreen"
    print(f"Python {sys.version.split()[0]} · {sys.executable}")
    check(Path(sys.executable).resolve().parent.name.lower() == "runtime", "Prüfung muss mit der eingebetteten Laufzeit laufen (runtime\\python.exe)")

    work = Path(tempfile.mkdtemp(prefix="pdf-tool-smoke-"))
    # Einstellungen der Prüfung in einem eigenen Ordner (vor dem Import von appstate setzen)
    os.environ["UE_DATA_DIR"] = str(work / "daten")
    sys.path.insert(0, str(app_dir))

    # 1. Module
    import importlib.util

    check(importlib.util.find_spec("tkinter") is None, "tkinter gehört nicht mehr zur Laufzeit (Qt-Oberfläche)")
    import PySide6
    from PySide6 import QtCore, QtGui, QtQml, QtQuick, QtQuickControls2, QtSvg, QtWidgets  # noqa: F401
    import numpy
    import openpyxl
    import pandas
    import PIL
    import pikepdf
    import pypdf
    import pypdfium2
    import reportlab
    import xlrd

    check(PySide6.__version__ == "6.11.2" and QtCore.qVersion() == "6.11.2", f"PySide6 {PySide6.__version__} / Qt {QtCore.qVersion()} statt 6.11.2")
    print(f"PySide6 {PySide6.__version__} (Qt {QtCore.qVersion()}) · numpy {numpy.__version__} · pandas {pandas.__version__} · openpyxl {openpyxl.__version__} · xlrd {xlrd.__version__} · reportlab {reportlab.Version} · Pillow {PIL.__version__}")
    print(f"pikepdf {pikepdf.__version__} (qpdf {pikepdf.__libqpdf_version__}) · pypdfium2 {pypdfium2.version.PYPDFIUM_INFO} (PDFium {pypdfium2.version.PDFIUM_INFO}) · pypdf {pypdf.__version__}")
    import appstate
    import engine
    import excelstyle
    import pdffonts
    import richtext
    from qtapp import application as qt_application
    from qtapp.contracts import comparison as qt_comparison
    from qtapp.contracts import tool as qt_contracts
    from qtapp import repair as qt_repair
    from tools import registry
    from tools.contract_overview import overview, preview
    from tools.contract_overview.batch import analyzer as batch_analyzer
    from tools.contract_overview.batch import models as batch_models
    from tools.contract_overview.batch import processor as batch_processor
    from tools.contract_overview.batch import resolver as batch_resolver
    from tools.contract_overview.customers import matching, migration, repository
    from tools.contract_overview.customers.texts import footer_of
    from tools.contract_overview.history import compare as history_compare
    from tools.contract_overview.history import models as history_models
    from tools.contract_overview.history import report as history_report
    from tools.contract_overview.history import repository as history_repository
    from tools.pdf_repair import engine as repair_engine
    from tools.pdf_repair import process as repair_process
    from tools.pdf_repair.models import Condition, Method, RepairStatus
    from tools.pdf_repair.recovery import lenient, rebuild, scanner

    check(lenient.available(), "pypdf fehlt in der Laufzeit (dritte Engine)")
    check(hasattr(qt_comparison, "ComparisonView") and qt_comparison.FIRST_SAVED.startswith("Erster Vertragsstand gespeichert"), "Vertragsvergleich fehlt")
    check(hasattr(rebuild, "write_classic") and hasattr(scanner, "scan"), "Rohrekonstruktion fehlt")

    check([tool.key for tool in registry.TOOLS] == ["contracts", "repair"], "Werkzeuge fehlen")
    check(hasattr(qt_contracts, "ContractsTool") and hasattr(qt_repair, "RepairTool"), "Werkzeuge der Oberfläche fehlen")
    check(registry.CONTRACTS.pages == ("create", "batch", "layout", "preview", "comparison", "customers"), "Ansichten von Vertragsübersichten fehlen")
    check(overview.contract_summary(5, 3) == "5 aktive Verträge · 3 inaktiv ausgeblendet" and overview.contract_summary(1) == "1 aktiver Vertrag", "Statuszeile der Excel-Prüfung")
    print(f"App {appstate.VERSION} · Schriften: {', '.join(pdffonts.available_families())} · Engines: {repair_engine.engine_name()}")
    if args.part == "ui":
        qml_check(app_dir, qt_application, full=args.ui)
        print("OK")
        return 0

    # 2. PDF mit Excel-Fettschrift und formatierter Fußzeile
    from datetime import datetime

    from openpyxl.styles import Font

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Vertrag-Nr.", "Beginnt am", "Abrechnungszyklus", "Netto [€]", "Zahlungsart", "Bemerkung", "Rechnungsempfänger Email", "Anwenderstatus"])
    sheet.append(["V-2", datetime(2022, 1, 1), "jährlich", 20.0, "Sofort", "Beta", "a@b.de", "Aktiv"])
    sheet.append(["V-1", datetime(2021, 1, 1), "monatlich", 10.0, "Lastschr", "Alpha", "a@b.de", "Aktiv"])
    sheet["D2"].font = Font(bold=True)
    excel = work / "probe.xlsx"
    book.save(excel)
    check(engine.pruefe_excel(excel)["fett"] == 1, "Fettschrift in der Excel nicht erkannt")
    style = richtext.FOOTER_STYLE
    text = "Kundennummer: {kd}"
    styles = [style] * 14 + [style.with_(bold=True, color="#B51F1F")] * 4
    fuss = richtext.RichText(text, styles, None, style, "center")
    logo = app_dir.parent / "assets" / "hott_logo_final.png"
    auftrag = engine.PdfAuftrag(excel=excel, logo=logo, kundennummer="4711", zielordner=work, fusszeile=fuss.text, fusszeile_format=fuss.to_dict())
    pdf = engine.erstelle_pdf(auftrag)
    data = pdf.read_bytes()
    check(data.startswith(b"%PDF") and len(data) > 5000, "PDF fehlt oder ist leer")
    check(b"Helvetica-Bold" in data, "keine fette Schrift in der PDF")
    print(f"PDF: {pdf.name} ({len(data)} Bytes)")
    check(excelstyle.SOURCE_ROW == "_source_excel_row", "Quellzeilen-Spalte unerwartet")

    # Kundenakte 2.0: Kundenhistorie wie aus 2.3 übernehmen (mit Sicherung), wiedererkennen, speichern
    import json

    akten = work / "kundenakten"
    akten.mkdir()
    alt = {
        "gesehen": "2.3.0",
        "kunden": [
            {"firmenname": "Müller & Söhne GmbH", "kundennummer": "0815", "rechnungsempfaenger": "Rechnung@Mueller.de", "fusszeile": fuss.text, "fusszeile_format": fuss.to_dict(), "kopfzeile": ""},
            {"firmenname": "Alt AG", "kundennummer": "1", "rechnungsempfaenger": "", "fusszeile": "", "kopfzeile": ""},
        ],
    }
    alt_datei = akten / "gui-config.json"
    alt_datei.write_text(json.dumps(alt, ensure_ascii=False, indent=2), encoding="utf-8")
    store, report = migration.open_store(akten, alt, alt_datei)
    check(report is not None and report.ok and report.migrated == 2, f"Übernahme der Kundenhistorie: {report}")
    check(report.backup is not None and report.backup.read_bytes() == alt_datei.read_bytes(), "Sicherung vor der Übernahme weicht ab")
    treffer = store.match([" RECHNUNG@mueller.de "])
    check(treffer.kind is matching.MatchKind.SINGLE, f"Wiedererkennung: {treffer.kind.value}")
    kunde = store.get(treffer.customer_id)
    check(footer_of(kunde) == fuss, "Formatierung der Fußzeile ging bei der Übernahme verloren")
    check(footer_of(next(k for k in store.all() if k.number == "1")) is None, "leere Fußzeile aus 2.3 wurde übernommen")
    check(store.match(["rechnung@mueller-gmbh.de"]).kind is matching.MatchKind.NONE, "unbekannte Adresse wurde zugeordnet")
    gespeichert = repository.CustomerStore.load(akten / repository.FILE_NAME)
    check(gespeichert.get(kunde.id) is not None and gespeichert.get(kunde.id).company == "Müller & Söhne GmbH", "Kundenakte nicht dauerhaft gespeichert")
    print(f"Kundenakte: {report.migrated} Kunden übernommen (Sicherung {report.backup.name}), Wiedererkennung: {treffer.kind.value}")

    # Vertragsvergleich: Stand der erstellten PDF speichern, unverändert nicht doppelt, Änderungen erkennen
    staende = history_repository.HistoryStore(work / "daten" / history_repository.FOLDER)
    records = history_models.records_from(auftrag.vertraege)
    check([r.contract_number for r in records] == ["V-1", "V-2"] and records[1].net_amount == "20.00", f"Vertragsstand der PDF: {records}")
    quelle = history_models.SnapshotSource(str(excel), history_repository.file_sha256(excel), str(pdf))
    erster, neu_angelegt = staende.record(kunde.id, records, quelle, kunde.label)
    zweiter, doppelt = staende.record(kunde.id, records, quelle, kunde.label)
    check(neu_angelegt and not doppelt and zweiter.id == erster.id and staende.count(kunde.id) == 1, "unveränderter Stand wurde doppelt gespeichert")
    from dataclasses import replace

    geaendert = (records[0], replace(records[1], net_amount="25.00", net_text="25,00 €"))
    vergleich = history_compare.compare(staende.latest(kunde.id), geaendert + (history_models.ContractRecord("V-3", "Softwarepflegevertrag", "Gamma", "2024-01-01", "jährlich", "5.00", "5,00 €", "Sofort"),))
    check(history_report.badges(vergleich) == "+1 neu · ~1 geändert" and len(vergleich.unchanged) == 1, f"Vergleich: {history_report.counts_text(vergleich)}")
    check(staende.snapshots("andere-kunden-id") == [], "Stände eines anderen Kunden sichtbar")
    print(f"Vertragsvergleich: 1 Stand gespeichert, {history_report.counts_text(vergleich)}")

    # Live-Vorschau: dieselbe PDF im Hintergrund erzeugen und die erste Seite mit PDFium zeichnen
    doc = preview.PreviewDocument.build(dict(excel=excel, logo=logo, kundennummer="4711", fusszeile=fuss.text, fusszeile_format=fuss.to_dict()))
    try:
        bild, breite, hoehe = doc.render(0, 1.0)
        check(doc.pages >= 1 and breite > 300 and hoehe > 300 and len(bild) > 1000, "Vorschau ist leer")
    finally:
        doc.close()
    check(not doc.folder.exists(), "Vorschau-Dateien wurden nicht gelöscht")
    print(f"Vorschau: {doc.pages} Seite(n), erste Seite {breite}×{hoehe} Pixel")

    # Stapel: zwei Listen – ein bekannter Kunde, ein unbekannter mit eigenen Angaben – mit derselben Engine
    stapel = work / "stapel"
    stapel.mkdir()
    for name, mail in (("bekannt.xlsx", "Rechnung@Mueller.de"), ("neu.xlsx", "info@neu.de")):
        liste = openpyxl.Workbook()
        blatt = liste.active
        blatt.append(["Vertrag-Nr.", "Beginnt am", "Abrechnungszyklus", "Netto [€]", "Zahlungsart", "Bemerkung", "Rechnungsempfänger Email", "Anwenderstatus"])
        blatt.append(["V-1", datetime(2021, 1, 1), "jährlich", 10.0, "Sofort", "Alpha", mail, "Aktiv"])
        blatt.append(["V-2", datetime(2022, 1, 1), "jährlich", 20.0, "Sofort", "Beta", mail, "Inaktiv"])
        liste.save(stapel / name)
    ausgabe = stapel / "ausgabe"
    settings = batch_models.BatchSettings(target_dir=str(ausgabe))
    defaults = batch_resolver.Defaults(
        dateiname=appstate.DEFAULT_DATEINAME, seitenformat="hoch", logo_breite="62", titel=appstate.DEFAULT_TITEL, untertitel=appstate.DEFAULT_UNTERTITEL,
        header=richtext.RichText.plain("", richtext.HEADER_STYLE), footer=appstate.default_footer_rich(), regeln=tuple(appstate.DEFAULT_REGELN), logo=str(logo),
    )
    items = {}
    for name in ("bekannt.xlsx", "neu.xlsx"):
        item = batch_models.BatchItem(str(stapel / name))
        item.analysis, item.stamp = batch_analyzer.analyze_file(item.path)
        items[item.id] = item
    bekannt, neu = items.values()
    check(bekannt.analysis.active == 1 and bekannt.analysis.inactive == 1, "Stapel: Prüfung der Excel")
    check(batch_resolver.resolve(bekannt, store, lambda _n: None, settings, defaults).ready, "Stapel: bekannter Kunde nicht bereit")
    offen = batch_resolver.resolve(neu, store, lambda _n: None, settings, defaults)
    check(batch_resolver.status_for(offen.issues) is batch_models.ItemStatus.NEEDS_INPUT, "Stapel: unbekannter Kunde sollte Angaben brauchen")
    neu.overrides = batch_models.Overrides(company="Neu GmbH", number="4712")
    ausgabe.mkdir()
    (ausgabe / "Vertragsuebersicht_Kd4712.pdf").write_bytes(b"%PDF-1.4 vorhanden")

    def job(item_id, created):
        res = batch_resolver.resolve(items[item_id], store, lambda _n: None, settings, defaults)
        return batch_processor.Job(item_id, res.company, items[item_id].path, items[item_id].stamp, batch_processor.identity_of(items[item_id].analysis), res.fields, res.folder, settings.conflict, created)

    fertig = {}
    runner = batch_processor.BatchRunner(list(items), job, lambda func, done, fail: done(func()), lambda *_a: None, lambda r: fertig.setdefault(r.item_id, r), lambda summary, _rest: fertig.setdefault("summary", summary))
    runner.start()
    summary = fertig["summary"]
    check(summary.created == 2 and summary.failed == 0, f"Stapel: {summary}")
    check((ausgabe / "Vertragsuebersicht_Kd0815.pdf").is_file(), "Stapel: PDF des bekannten Kunden fehlt")
    check((ausgabe / "Vertragsuebersicht_Kd4712.pdf").read_bytes() == b"%PDF-1.4 vorhanden", "Stapel: vorhandene PDF wurde überschrieben")
    check((ausgabe / "Vertragsuebersicht_Kd4712_2.pdf").is_file(), "Stapel: nummerierte PDF fehlt")
    check(not any(path.suffix == ".tmp" for path in ausgabe.iterdir()), "Stapel: temporäre Dateien blieben zurück")
    print(f"Stapel: {summary.created} Übersichten erstellt ({', '.join(sorted(p.name for p in ausgabe.iterdir()))})")

    # 3. PDF reparieren – im eigenen Arbeitsprozess wie in der App
    import re

    folder = work / "Übersichten März"
    folder.mkdir()
    damaged = folder / "Vertrag beschädigt.pdf"
    damaged.write_bytes(re.sub(rb"startxref\s+(\d+)", lambda m: b"startxref\n" + str(int(m.group(1)) + 97).encode(), data))
    before = damaged.read_bytes()
    job = repair_process.Job("analyze", {"path": str(damaged), "password": None})
    events = job.wait(120)
    results = [event[1] for event in events if event[0] == "result"]
    check(bool(results), f"Analyse ohne Ergebnis: {events[-1:]}")
    analysis = results[0]
    check(analysis.condition is Condition.REPAIRABLE, f"Analyse: {analysis.condition.value}")
    job = repair_process.Job("repair", {"path": str(damaged), "password": None, "mode": "auto", "sha256": analysis.sha256})
    events = job.wait(120)
    results = [event[1] for event in events if event[0] == "result"]
    check(bool(results), f"Reparatur ohne Ergebnis: {events[-1:]}")
    result = results[0]
    check(result.status is RepairStatus.REPAIRED, f"Reparatur: {result.status.value} {result.error}")
    output = repair_process.deliver(Path(result.output_path), damaged)
    check(output == folder / "Vertrag beschädigt_repariert.pdf", f"Ausgabe: {output}")
    # Mehrere PDFs: Liste ohne Doppelte, Name ohne »_repariert« – das Original wird nie überschrieben
    from tools.pdf_repair.batch import RepairBatch

    liste = RepairBatch()
    check(len(liste.add([damaged]).added) == 1 and not liste.add([damaged]).added, "Liste: dieselbe Datei doppelt")
    liste.plan(lambda item: item.path.parent, False, "_repariert")
    eintrag = liste.items[0]
    check(eintrag.planned == folder / "Vertrag beschädigt (1).pdf", f"Geplanter Name ohne Zusatz: {eintrag.planned}")
    ohne = repair_process.deliver(Path(result.output_path), damaged, folder, base=eintrag.planned_base, start=eintrag.planned_number)
    check(ohne == eintrag.planned, f"Ausgabe ohne Zusatz: {ohne}")
    job.cleanup()
    check(damaged.read_bytes() == before, "Original wurde verändert")
    with pikepdf.open(output) as repaired:
        check(len(repaired.pages) == result.pages_before, "Seitenzahl der reparierten PDF stimmt nicht")
    with pypdfium2.PdfDocument(str(output)) as second:
        check(len(second) == result.pages_before, "PDFium liest die reparierte PDF nicht")
    print(f"Reparatur: {output.name} ({result.method.value}, {result.pages_after} Seiten, {output.stat().st_size} Bytes)")

    # Erweiterte Wiederherstellung: klassische PDF 1.4 ohne xref, Trailer und %%EOF, Seitenbaum zeigt ins Leere
    objs = {1: b"<< /Type /Catalog /Pages 2 0 R >>", 2: b"<< /Type /Pages /Kids [ 99 0 R ] /Count 1 >>", 3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"}
    for index in range(3):
        inhalt = f"BT /F1 24 Tf 72 700 Td (Seite {index + 1} von 3) Tj ET".encode()
        objs[5 + 2 * index] = b"<< /Length %d >>\nstream\n" % len(inhalt) + inhalt + b"\nendstream"
        objs[4 + 2 * index] = b"<< /Type /Page /Parent 2 0 R /Contents %d 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 3 0 R >> >> >>" % (5 + 2 * index)
    roh = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    for number in sorted(objs):
        roh += b"%d 0 obj\n" % number + objs[number] + b"\nendobj\n"
    zerstoert = folder / "Struktur zerstört.pdf"
    zerstoert.write_bytes(bytes(roh) + b"xref\n0 12\ngarbage 00000 f\n")
    vorher = zerstoert.read_bytes()
    job = repair_process.Job("analyze", {"path": str(zerstoert), "password": None})
    results = [event[1] for event in job.wait(120) if event[0] == "result"]
    check(bool(results) and results[0].condition is Condition.RAW_RECOVERABLE, f"Analyse (erweitert): {results[0].condition.value if results else 'ohne Ergebnis'}")
    check(results[0].raw is not None and results[0].raw.objects == 9 and results[0].raw.pages == 3, f"Rohanalyse: {results[0].raw}")
    job = repair_process.Job("repair", {"path": str(zerstoert), "password": None, "mode": "auto", "sha256": results[0].sha256})
    results = [event[1] for event in job.wait(120) if event[0] == "result"]
    check(bool(results), "Rekonstruktion ohne Ergebnis")
    result = results[0]
    check(result.status is RepairStatus.REPAIRED and result.method is Method.PAGE_TREE_REBUILD and result.pages_after == 3, f"Rekonstruktion: {result.status.value} {result.method} {result.pages_after} {result.error}")
    output = repair_process.deliver(Path(result.output_path), zerstoert)
    job.cleanup()
    check(zerstoert.read_bytes() == vorher, "Original wurde bei der Rekonstruktion verändert")
    texte = [(seite.extract_text() or "").strip() for seite in pypdf.PdfReader(str(output)).pages]
    check(texte == ["Seite 1 von 3", "Seite 2 von 3", "Seite 3 von 3"], f"Text der rekonstruierten PDF: {texte}")
    with pikepdf.open(output, attempt_recovery=False) as rebuilt:
        check(rebuilt.get_warnings() == [] and len(rebuilt.pages) == 3, "rekonstruierte PDF öffnet nicht ohne Wiederherstellung")
    print(f"Erweiterte Wiederherstellung: {output.name} ({result.method.value}, {result.pages_after} Seiten)")

    if args.part == "runtime":
        print("OK")
        return 0
    # 4. Oberfläche
    qml_check(app_dir, qt_application, full=args.ui)

    print("OK")
    return 0


def qml_check(app_dir: Path, qt_application, full: bool) -> None:
    """QML aus der eingebauten Ressource laden, ohne Meldungen der QML-Engine; dann Programmstart."""
    source, _import_path = qt_application.qml_source()
    if (app_dir / "qml_rc.py").is_file():
        check(source.toString() == "qrc:/qml/Main.qml", f"Oberfläche nicht aus der Ressource: {source.toString()}")
        check(not (app_dir / "qml").exists(), "lose QML-Dateien im Programmordner")
    ui_probe(qt_application, full=full)


def ui_probe(qt_application, full: bool) -> None:
    """Hauptfenster wie per Verknüpfung starten; mit ``full`` Werkzeuge, Ansichten und Kundenakte prüfen."""
    import json

    from PySide6.QtCore import QTimer, qInstallMessageHandler

    from qtapp import dialogs

    dialogs.AUTO_ANSWER = "primary"  # z. B. »Neu in Version« schließt sich selbst
    qInstallMessageHandler(qt_application._message_handler)
    started = time.monotonic()
    qt = qt_application.create_application([])
    runtime = qt_application.Runtime(qt_application.load_config())
    engine_qml = qt_application.create_engine(runtime)
    window = qt_application.show_window(runtime, engine_qml)
    app = runtime.app
    shown: dict = {}

    def step(actions: list) -> None:
        if not actions:
            return
        action = actions.pop(0)
        try:
            action()
        except Exception:  # noqa: BLE001 - Fehler als Ergebnis festhalten
            shown["error"] = traceback.format_exc()
            actions.clear()
            qt.quit()
            return
        QTimer.singleShot(250, lambda: step(actions))

    def first() -> None:
        shown["visible"] = window.isVisible()
        shown["page"] = app.currentPage
        shown["title"] = window.title()
        shown["ready"] = app.ready

    def tools() -> None:
        app.openTool("repair")
        shown["repair"] = app.currentPage
        app.openTool("contracts")
        shown["contracts"] = app.currentPage

    def views() -> None:
        for view in ("batch", "preview"):
            app.navigate(view)
            shown[view] = app.currentPage

    def records_off() -> None:
        shown["records"] = runtime.settings.customerRecords
        shown["customers_blocked"] = "customers" in app.unavailablePages
        app.navigate("customers")
        shown["customers_off"] = app.currentPage

    def records_on() -> None:
        runtime.settings.setCustomerRecords(True)

    def open_customers() -> None:
        app.navigate("customers")
        shown["customers"] = app.currentPage
        runtime.settings.setCustomerRecords(False)

    def finish() -> None:
        shown["customers_after"] = app.currentPage
        shown["messages"] = list(qt_application.MESSAGES)
        shown["closed"] = app.requestClose()
        window.close()
        qt.quit()

    actions = [first, tools, views, records_off, records_on, open_customers, finish] if full else [first, finish]
    QTimer.singleShot(2500, lambda: step(actions))
    QTimer.singleShot(60000, qt.quit)  # Sicherheitsnetz
    qt.exec()
    qt_application.finish_incubation(engine_qml)  # wie beim Beenden der App
    del engine_qml
    check("error" not in shown, f"Fehler beim Programmstart:\n{shown.get('error')}")
    check(shown.get("visible") is True and shown.get("ready") is True, "Hauptfenster wurde nicht angezeigt")
    check(shown.get("page") == "home", "Startseite fehlt")
    check(shown.get("title") == "PDF Tool", f"Fenstertitel: {shown.get('title')}")
    check(not shown.get("messages"), "Meldungen der QML-Engine:\n" + "\n".join(shown.get("messages") or []))
    check(shown.get("closed") is True, "App ließ sich nicht schließen")
    if full:
        check(shown.get("repair") == "repair" and shown.get("contracts") == "create", "Werkzeuge lassen sich nicht öffnen")
        check(shown.get("preview") == "preview" and shown.get("batch") == "batch", "Ansichten »Stapel« und »Vorschau« lassen sich nicht öffnen")
        check(shown.get("records") is False and shown.get("customers_blocked") is True and shown.get("customers_off") == "preview", "Kundenakte ist nicht standardmäßig aus bzw. »Kunden« ohne Kundenakte erreichbar")
        check(shown.get("customers") == "customers" and shown.get("customers_after") != "customers", "Kundenakte lässt sich nicht ohne Neustart ein- und ausschalten")
    config = Path(os.environ["UE_DATA_DIR"]) / "gui-config.json"
    check(config.is_file() and isinstance(json.loads(config.read_text(encoding="utf-8")), dict), "Einstellungen wurden beim Beenden nicht gespeichert")
    print(f"Oberfläche: Fenster sichtbar, {'Werkzeuge und Ansichten geprüft, ' if full else ''}beendet nach {time.monotonic() - started:.1f} s")


if __name__ == "__main__":
    try:
        sys.exit(main())
    except SystemExit:
        raise
    except BaseException:
        traceback.print_exc()
        sys.exit(1)
