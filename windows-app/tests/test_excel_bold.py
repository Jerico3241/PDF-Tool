"""Excel-Fettschrift in der PDF (Version 2.2): Tests A–I und Absicherung der Zuordnung.

Geprüft wird in der erzeugten PDF, mit welcher Schrift jede Tabellenzelle gesetzt
ist: Helvetica (normal) oder Helvetica-Bold (aus Excel übernommen).
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest

import appstate
import engine
import excelstyle

pypdf = pytest.importorskip("pypdf")

SPALTEN = ["Vertrag-Nr.", "Beginnt am", "Abrechnungszyklus", "Netto [€]", "Zahlungsart", "Bemerkung", "Rechnungsempfänger Email", "Anwenderstatus"]
# Spaltenbuchstaben dieser Testdateien (nur für das Erzeugen der Dateien – die App erkennt Spalten über die Überschrift)
SPALTE = {"nr": "A", "beginn": "B", "zyklus": "C", "netto": "D", "zahlung": "E", "bemerkung": "F"}
PDF_ZELLEN = ("art", "nr", "bemerkung", "beginn", "zyklus", "netto", "zahlung")


def zeile(nr: str, jahr: int, netto: float, bemerkung: str, status: str = "Aktiv", zahlung: str = "Sofort", zyklus: str = "jährlich") -> list:
    return [nr, datetime(jahr, 1, 1), zyklus, netto, zahlung, bemerkung, "rechnung@muster.de", status]


def schreibe_xlsx(pfad: Path, zeilen: list[list], fett: dict[int, tuple[str, ...]] | None = None, titelzeilen: int = 0, spalten: list[str] = SPALTEN) -> Path:
    """``fett``: Datenzeile (0-basiert) → fette Spalten (Schlüssel aus SPALTE)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    sheet = book.active
    for index in range(titelzeilen):
        sheet.append(["Vertragsliste Export" if index == 0 else None])
    sheet.append(spalten)
    kopf = titelzeilen + 1
    for werte in zeilen:
        sheet.append(werte)
    for datenzeile, keys in (fett or {}).items():
        for key in keys:
            buchstaben = [SPALTE[key]] if key in SPALTE else [chr(ord("A") + i) for i in range(len(spalten))]
            for b in buchstaben:
                sheet[f"{b}{kopf + 1 + datenzeile}"].font = Font(bold=True)
    book.save(pfad)
    return pfad


def pdf_zellen(pfad: Path) -> dict[str, set[str]]:
    """Text jeder Tabellenzelle → verwendete Schriften (für eindeutige Texte)."""
    schriften: dict[str, set[str]] = {}
    for text, font, _x, _y in pdf_textstuecke(pfad):
        schriften.setdefault(text, set()).add(font)
    return schriften


def pdf_textstuecke(pfad: Path) -> list[tuple[str, str, float, float]]:
    """Alle Textstücke der ersten Seite: (Text, Schrift, x, y)."""
    stuecke: list[tuple[str, str, float, float]] = []

    def visit(text, cm, tm, font, _size):
        text = " ".join(str(text).split())
        if text and font:
            x = tm[4] * cm[0] + cm[4]
            y = tm[5] * cm[3] + cm[5]
            stuecke.append((text, str(font.get("/BaseFont", "")).lstrip("/"), round(x, 1), round(y, 1)))

    pypdf.PdfReader(str(pfad)).pages[0].extract_text(visitor_text=visit)
    return stuecke


def pdf_zeile(pfad_oder_stuecke, nr: str) -> dict[str, tuple[str, str]]:
    """Tabellenzeile eines Vertrags: Spalte → (Text der ersten Zeile, Schrift).

    Alle Zellen einer Zeile beginnen oben auf derselben Grundlinie (VALIGN TOP);
    von links nach rechts folgen Art, Vertrag-Nr., Beschreibung, Beginn, Zyklus, Netto, Zahlungsart.
    """
    stuecke = pdf_textstuecke(pfad_oder_stuecke) if isinstance(pfad_oder_stuecke, Path) else pfad_oder_stuecke
    y = next(sy for text, _f, _x, sy in stuecke if text == nr)
    reihe = sorted((sx, text, font) for text, font, sx, sy in stuecke if abs(sy - y) < 0.6)
    assert len(reihe) == len(PDF_ZELLEN), f"Zeile {nr}: {reihe}"
    return {key: (text, font) for key, (_x, text, font) in zip(PDF_ZELLEN, reihe)}


def erstelle(excel: Path, ziel: Path) -> Path:
    return engine.erstelle_pdf(engine.PdfAuftrag(excel=excel, logo=appstate.DEFAULT_LOGO, kundennummer="10042", zielordner=ziel))


def ist_fett(schriften: dict[str, set[str]], text: str) -> bool:
    fonts = schriften.get(text)
    assert fonts, f"{text!r} nicht in der PDF gefunden"
    assert len(fonts) == 1, f"{text!r} uneinheitlich gesetzt: {fonts}"
    return next(iter(fonts)) == "Helvetica-Bold"


def zellen_von(nr: str, jahr: int, netto: str, bemerkung: str, zahlung: str = "Sofort", zyklus: str = "jährlich") -> dict[str, str]:
    art = "Supportvertrag" if "hotline" in bemerkung.lower() else "Software-"
    return {"art": art, "nr": nr, "bemerkung": bemerkung, "beginn": f"01.01.{jahr}", "zyklus": zyklus, "netto": netto, "zahlung": zahlung}


def pruefe_zeile(quelle, zellen: dict[str, str], fett: set[str]) -> None:
    """``quelle``: PDF-Pfad oder Textstücke. Prüft Text und Schriftgewicht jeder Zelle der Zeile."""
    reihe = pdf_zeile(quelle, zellen["nr"])
    for key in PDF_ZELLEN:
        text, font = reihe[key]
        # Erste Zeile der Zelle: Lange Wörter können in schmalen Spalten umbrechen (fett ist etwas breiter).
        assert text and zellen[key].startswith(text), f"{key}: {text!r} statt {zellen[key]!r}"
        erwartet = "Helvetica-Bold" if key in fett else "Helvetica"
        assert font == erwartet, f"{key} ({text!r}) sollte {'fett' if key in fett else 'normal'} sein, ist {font}"


# --- A–E: zellweise Übernahme ---------------------------------------------------------------


def test_a_normale_zeile_bleibt_normal(tmp_path: Path) -> None:
    excel = schreibe_xlsx(tmp_path / "a.xlsx", [zeile("V-100", 2021, 10.0, "Warenwirtschaft Alpha")])
    schriften = pdf_textstuecke(erstelle(excel, tmp_path))
    pruefe_zeile(schriften, zellen_von("V-100", 2021, "10,00", "Warenwirtschaft Alpha"), set())


def test_b_nur_beschreibung_fett(tmp_path: Path) -> None:
    excel = schreibe_xlsx(tmp_path / "b.xlsx", [zeile("V-100", 2021, 10.0, "Warenwirtschaft Alpha")], fett={0: ("bemerkung",)})
    schriften = pdf_textstuecke(erstelle(excel, tmp_path))
    # »Art« wird aus der Bemerkung abgeleitet und übernimmt deren Formatierung.
    pruefe_zeile(schriften, zellen_von("V-100", 2021, "10,00", "Warenwirtschaft Alpha"), {"bemerkung", "art"})


def test_c_nur_netto_fett(tmp_path: Path) -> None:
    excel = schreibe_xlsx(tmp_path / "c.xlsx", [zeile("V-100", 2021, 1234.5, "Warenwirtschaft Alpha")], fett={0: ("netto",)})
    schriften = pdf_textstuecke(erstelle(excel, tmp_path))
    pruefe_zeile(schriften, zellen_von("V-100", 2021, "1.234,50", "Warenwirtschaft Alpha"), {"netto"})


def test_d_mehrere_einzelne_zellen(tmp_path: Path) -> None:
    excel = schreibe_xlsx(tmp_path / "d.xlsx", [zeile("V-100", 2021, 10.0, "Warenwirtschaft Alpha", zahlung="Lastschr")], fett={0: ("nr", "zahlung", "zyklus")})
    schriften = pdf_textstuecke(erstelle(excel, tmp_path))
    pruefe_zeile(schriften, zellen_von("V-100", 2021, "10,00", "Warenwirtschaft Alpha", zahlung="Lastschrift"), {"nr", "zahlung", "zyklus"})


def test_e_ganze_zeile_fett(tmp_path: Path) -> None:
    excel = schreibe_xlsx(tmp_path / "e.xlsx", [zeile("V-100", 2021, 10.0, "Hotline Premium"), zeile("V-200", 2022, 20.0, "Warenwirtschaft Beta")], fett={0: ("alle",)})
    schriften = pdf_textstuecke(erstelle(excel, tmp_path))
    pruefe_zeile(schriften, zellen_von("V-100", 2021, "10,00", "Hotline Premium"), set(PDF_ZELLEN))
    pruefe_zeile(schriften, zellen_von("V-200", 2022, "20,00", "Warenwirtschaft Beta"), set())


# --- F–G: Sortierung und Filter zerstören die Zuordnung nicht ---------------------------------


def test_f_sortierung_nach_vertragsbeginn(tmp_path: Path) -> None:
    # In der Excel: 2023, 2021, 2022 – in der PDF sortiert nach Beginn: 2021, 2022, 2023.
    zeilen = [zeile("V-2023", 2023, 30.0, "Gamma"), zeile("V-2021", 2021, 10.0, "Alpha"), zeile("V-2022", 2022, 20.0, "Beta")]
    excel = schreibe_xlsx(tmp_path / "f.xlsx", zeilen, fett={0: ("netto",), 2: ("bemerkung",)})
    pfad = erstelle(excel, tmp_path)
    schriften = pdf_textstuecke(pfad)
    pruefe_zeile(schriften, zellen_von("V-2023", 2023, "30,00", "Gamma"), {"netto"})
    pruefe_zeile(schriften, zellen_von("V-2022", 2022, "20,00", "Beta"), {"bemerkung", "art"})
    pruefe_zeile(schriften, zellen_von("V-2021", 2021, "10,00", "Alpha"), set())
    text = pypdf.PdfReader(str(pfad)).pages[0].extract_text()
    assert text.index("V-2021") < text.index("V-2022") < text.index("V-2023")


def test_g_inaktive_zeilen_zwischen_aktiven(tmp_path: Path) -> None:
    zeilen = [
        zeile("V-1", 2021, 10.0, "Alpha"),
        zeile("V-X1", 2020, 99.0, "Alt eins", status="Inaktiv"),
        zeile("V-2", 2022, 20.0, "Beta"),
        zeile("V-X2", 2019, 98.0, "Alt zwei", status="inaktiv"),
        zeile("V-3", 2023, 30.0, "Gamma"),
    ]
    # Fett: inaktive Zeilen komplett, bei den aktiven jeweils eine andere Zelle
    fett = {1: ("alle",), 3: ("alle",), 0: ("nr",), 2: ("zahlung",), 4: ("beginn",)}
    excel = schreibe_xlsx(tmp_path / "g.xlsx", zeilen, fett=fett)
    schriften = pdf_textstuecke(erstelle(excel, tmp_path))
    assert not any(text in ("V-X1", "V-X2") for text, *_rest in schriften)
    pruefe_zeile(schriften, zellen_von("V-1", 2021, "10,00", "Alpha"), {"nr"})
    # Zahlungsart »Sofort« kommt in mehreren Zeilen vor – eigener Wert für die fette Zelle
    pdf = erstelle(schreibe_xlsx(tmp_path / "g2.xlsx", [zeile("V-1", 2021, 10.0, "Alpha"), zeile("V-X", 2020, 1.0, "Alt", status="Inaktiv"), zeile("V-2", 2022, 20.0, "Beta", zahlung="Überweisung")], fett={1: ("alle",), 2: ("zahlung",)}), tmp_path / "g2")
    pruefe_zeile(pdf, zellen_von("V-2", 2022, "20,00", "Beta", zahlung="Überweisung"), {"zahlung"})
    pruefe_zeile(pdf, zellen_von("V-1", 2021, "10,00", "Alpha"), set())
    pruefe_zeile(schriften, zellen_von("V-3", 2023, "30,00", "Gamma"), {"beginn"})


def test_leerzeilen_und_titelzeilen_vor_der_tabelle(tmp_path: Path) -> None:
    """Überschrift nicht in Zeile 1 und Leerzeilen zwischen den Verträgen."""
    zeilen = [zeile("V-1", 2021, 10.0, "Alpha"), [None] * len(SPALTEN), zeile("V-2", 2022, 20.0, "Beta"), [None] * len(SPALTEN), zeile("V-3", 2023, 30.0, "Gamma")]
    excel = schreibe_xlsx(tmp_path / "t.xlsx", zeilen, fett={2: ("netto",), 4: ("nr",)}, titelzeilen=2)
    tabelle = engine.lies_tabelle(excel)
    assert tabelle.kopfzeile == 3
    nummern = dict(zip(tabelle.df[excelstyle.SOURCE_ROW], tabelle.df["Vertrag-Nr."]))
    assert nummern[4] == "V-1" and nummern[6] == "V-2" and nummern[8] == "V-3"
    schriften = pdf_textstuecke(erstelle(excel, tmp_path))
    pruefe_zeile(schriften, zellen_von("V-2", 2022, "20,00", "Beta"), {"netto"})
    pruefe_zeile(schriften, zellen_von("V-3", 2023, "30,00", "Gamma"), {"nr"})
    pruefe_zeile(schriften, zellen_von("V-1", 2021, "10,00", "Alpha"), set())


def test_spalten_werden_ueber_die_ueberschrift_erkannt(tmp_path: Path) -> None:
    """Andere Spaltenreihenfolge und -namen (Beschreibung, Beginn, Netto): keine festen Buchstaben."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    sheet = book.active
    sheet.append(["Anwenderstatus", "Beschreibung", "Netto", "Zahlungsart", "Beginn", "Vertragsnummer", "Abrechnungszyklus"])
    sheet.append(["Aktiv", "Warenwirtschaft", 15.0, "Sofort", datetime(2021, 5, 1), "K-7", "monatlich"])
    sheet["B2"].font = Font(bold=True)  # Beschreibung
    sheet["F2"].font = Font(bold=True)  # Vertragsnummer
    book.save(tmp_path / "s.xlsx")
    schriften = pdf_textstuecke(erstelle(tmp_path / "s.xlsx", tmp_path))
    pruefe_zeile(schriften, zellen_von("K-7", 2021, "15,00", "Warenwirtschaft", zyklus="monatlich") | {"beginn": "01.05.2021"}, {"nr", "bemerkung", "art"})


# --- H–I: Dateiformate ----------------------------------------------------------------------------


def test_h_xlsx_mapping_und_pruefung(tmp_path: Path) -> None:
    excel = schreibe_xlsx(tmp_path / "h.xlsx", [zeile("V-1", 2021, 10.0, "Alpha"), zeile("V-2", 2022, 20.0, "Beta")], fett={1: ("netto", "nr")})
    assert excelstyle.file_kind(excel) == "xlsx"
    ergebnis = engine.pruefe_excel(excel)
    assert ergebnis["fett"] == 2
    schriften = pdf_textstuecke(erstelle(excel, tmp_path))
    pruefe_zeile(schriften, zellen_von("V-2", 2022, "20,00", "Beta"), {"netto", "nr"})


def test_i_xls_mit_formatierung(tmp_path: Path) -> None:
    xlwt = pytest.importorskip("xlwt")
    book = xlwt.Workbook()
    sheet = book.add_sheet("Liste")
    fett = xlwt.easyxf("font: bold on")
    for col, name in enumerate(SPALTEN):
        sheet.write(0, col, name)
    daten = [("V-1", "2021-01-01", "jährlich", 10.0, "Sofort", "Alpha", "a@b.de", "Aktiv"), ("V-X", "2020-01-01", "jährlich", 5.0, "Sofort", "Alt", "a@b.de", "Inaktiv"), ("V-2", "2022-01-01", "jährlich", 20.0, "Sofort", "Beta", "a@b.de", "Aktiv")]
    for row, werte in enumerate(daten, start=1):
        for col, wert in enumerate(werte):
            if (row == 3 and col in (0, 3)) or row == 2:
                sheet.write(row, col, wert, fett)
            else:
                sheet.write(row, col, wert)
    pfad = tmp_path / "i.xls"
    book.save(str(pfad))
    assert excelstyle.file_kind(pfad) == "xls"
    schriften = pdf_textstuecke(erstelle(pfad, tmp_path))
    pruefe_zeile(schriften, zellen_von("V-2", 2022, "20,00", "Beta"), {"nr", "netto"})
    pruefe_zeile(schriften, zellen_von("V-1", 2021, "10,00", "Alpha"), set())


def test_i_xls_ohne_lesbare_formatierung_sicherer_fallback(tmp_path: Path, monkeypatch) -> None:
    """Liefert eine XLS keine Stilinformationen, entsteht die PDF trotzdem – normal formatiert."""
    xlwt = pytest.importorskip("xlwt")
    book = xlwt.Workbook()
    sheet = book.add_sheet("Liste")
    for col, name in enumerate(SPALTEN):
        sheet.write(0, col, name)
    for col, wert in enumerate(("V-1", "2021-01-01", "jährlich", 10.0, "Sofort", "Alpha", "a@b.de", "Aktiv")):
        sheet.write(1, col, wert, xlwt.easyxf("font: bold on"))
    pfad = tmp_path / "exotisch.xls"
    book.save(str(pfad))

    def kaputt(*_args, **_kwargs):
        raise NotImplementedError("formatting_info not supported")

    monkeypatch.setattr(excelstyle, "_read_xlrd", kaputt)
    stile = excelstyle.read_styles(pfad, [2], [1])
    assert not stile.available and "Formatierung" in stile.reason
    schriften = pdf_textstuecke(erstelle(pfad, tmp_path))
    pruefe_zeile(schriften, zellen_von("V-1", 2021, "10,00", "Alpha"), set())


def test_xls_endung_mit_xlsx_inhalt(tmp_path: Path) -> None:
    """Exporte heißen manchmal .xls, sind aber .xlsx: Erkennung am Inhalt wie bei pandas."""
    echt = schreibe_xlsx(tmp_path / "x.xlsx", [zeile("V-1", 2021, 10.0, "Alpha")], fett={0: ("nr",)})
    falsch = tmp_path / "export.xls"
    falsch.write_bytes(echt.read_bytes())
    assert excelstyle.file_kind(falsch) == "xlsx"


# --- Absicherung --------------------------------------------------------------------------------


def test_zuordnung_wird_an_vertragsnummern_geprueft(tmp_path: Path) -> None:
    excel = schreibe_xlsx(tmp_path / "p.xlsx", [zeile("V-1", 2021, 10.0, "Alpha")], fett={0: ("netto",)})
    richtig = excelstyle.read_styles(excel, [2], [1, 4], expected={(2, 1): "V-1"})
    assert richtig.available and richtig.bold(2, 4) and not richtig.bold(2, 1)
    falsch = excelstyle.read_styles(excel, [2], [1, 4], expected={(2, 1): "V-2"})
    assert not falsch.available and not falsch.bold(2, 4)


def test_quellzeile_bleibt_nach_filter_und_sortierung(tmp_path: Path) -> None:
    zeilen = [zeile("V-3", 2023, 30.0, "C"), zeile("V-X", 2020, 1.0, "X", status="Inaktiv"), zeile("V-1", 2021, 10.0, "A")]
    excel = schreibe_xlsx(tmp_path / "q.xlsx", zeilen)
    tabelle = engine.lies_tabelle(excel)
    df = tabelle.df.dropna(subset=["Vertrag-Nr."])
    df, _ = engine.nur_aktive(df)
    df = df.sort_values("Beginnt am", kind="stable").reset_index(drop=True)
    assert list(df["Vertrag-Nr."]) == ["V-1", "V-3"]
    assert list(df[excelstyle.SOURCE_ROW]) == [4, 2]


def test_nur_die_ganze_zelle_zaehlt(tmp_path: Path) -> None:
    """Teilweise fette Zellen (Rich Text in Excel) werden in 2.2 nicht zeichenweise übernommen."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    book = Workbook()
    sheet = book.active
    sheet.append(SPALTEN)
    sheet.append(zeile("V-1", 2021, 10.0, "Alpha Beta"))
    sheet["F2"] = CellRichText(["Alpha ", TextBlock(InlineFont(b=True), "Beta")])
    book.save(tmp_path / "r.xlsx")
    assert load_workbook(tmp_path / "r.xlsx")["Sheet"]["F2"].value is not None
    schriften = pdf_zellen(erstelle(tmp_path / "r.xlsx", tmp_path))
    assert not ist_fett(schriften, "Alpha Beta")


def test_pdf_design_bleibt_gleich_fuer_fette_zellen(tmp_path: Path) -> None:
    """Fett ändert nur das Schriftgewicht: gleiche Größe und Zeilenposition wie normale Zellen."""
    excel = schreibe_xlsx(tmp_path / "z.xlsx", [zeile("V-1", 2021, 10.0, "Alpha")], fett={0: ("netto",)})
    pfad = erstelle(excel, tmp_path)
    groessen: dict[str, tuple[float, float]] = {}

    def visit(text, cm, tm, font, size):
        text = " ".join(str(text).split())
        if text in ("10,00", "Sofort"):
            groessen[text] = (round(size * tm[0], 2), round(tm[5] * cm[3] + cm[5], 1))

    pypdf.PdfReader(str(pfad)).pages[0].extract_text(visitor_text=visit)
    assert groessen["10,00"] == groessen["Sofort"]
