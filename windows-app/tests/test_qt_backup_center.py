"""Qt-Oberfläche (5.0 Beta 7): »Backup-Center« – Fassungen beim Speichern, automatische Sicherung ungespeicherter
Stände, Wiederherstellen und Einstellungen.

Die Abläufe laufen über den Controller (``Versions``) wie ein Klick in QML; geprüft wird, dass die Seite ohne
QML-Meldungen lädt, das Original nur nach Rückfrage ersetzt wird und ein geöffnetes Dokument nie überschrieben wird.
"""

from __future__ import annotations

import hashlib
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent))

import editorsamples as samples
from conftest import neustart, pump, wait_until

from tools.pdf_editor import fonts


@pytest.fixture(autouse=True)
def substitute_fonts(monkeypatch, tmp_path: Path):
    monkeypatch.setenv("PDFTOOL_FONT_DIRS", str(samples.vera_path().parent))
    monkeypatch.setenv("NOVADESK_FASSUNGEN", str(tmp_path / "Fassungen"))
    fonts.find_system_font.cache_clear()
    yield
    fonts.find_system_font.cache_clear()


def controller(h):
    return h.runtime.versions.controller


def reader(h):
    return h.runtime.reader.controller


def digest(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def settle(h, timeout: float = 60.0) -> None:
    r = reader(h)
    assert wait_until(lambda: r.engine.idle() and r.opening == 0 and (r.current is None or not r.current.busy), timeout)
    pump(0.15)


def open_pdf(h, path: Path):
    r = reader(h)
    r.open_paths([str(path)])
    assert wait_until(lambda: r.current is not None and r.current.path and Path(r.current.path) == path, 30)
    settle(h)
    return r.current


def refreshed(h):
    c = controller(h)
    c.refresh()
    assert wait_until(lambda: not c.loading, 20)
    pump(0.1)
    return c


def test_page_empty_then_lists_versions_after_save(ui_app, tmp_path: Path) -> None:
    h = ui_app
    h.app.openTool("versions")
    pump(0.4)
    assert h.app.currentPage == "versions" and h.item("versionsPage") is not None
    assert h.item("versionsEmpty").isVisible()
    path = samples.standard_text(tmp_path / "Brief.pdf")
    before = digest(path)
    doc = open_pdf(h, path)
    doc.rotatePages([0], 90)
    settle(h)
    doc.saveDocument()
    settle(h)
    assert not doc.dirty and digest(path) != before
    h.app.openTool("versions")
    c = refreshed(h)
    assert [row["name"] for row in c.documents] == ["Brief.pdf"] and c.documents[0]["exists"]
    assert c.selected == c.documents[0]["key"] and len(c.versionList) == 1
    version = c.versionList[0]
    assert version["kindText"] == "Vor dem Speichern" and version["newest"] and digest(Path(version["path"])) == before
    assert h.item("versionsDocument_0").isVisible() and h.item("versionsVersion_0").isVisible()
    assert c.totalText.startswith("1 Dokument · ")
    # Als Kopie wiederherstellen: neue Datei neben dem Original, das Original bleibt
    saved = digest(path)
    c.restoreCopy(version["path"])
    assert wait_until(lambda: len(list(tmp_path.glob("Brief (Fassung *).pdf"))) == 1, 20)
    copy = next(tmp_path.glob("Brief (Fassung *).pdf"))
    assert digest(copy) == before and digest(path) == saved
    # Öffnen zeigt eine Kopie im Reader (nicht die Ablage selbst)
    c.openVersion(version["path"])
    assert wait_until(lambda: reader(h).current is not None and "(Fassung " in reader(h).current.name, 30)
    settle(h)
    assert Path(reader(h).current.path).parent != Path(version["path"]).parent
    assert not h.messages()


def test_replace_original_refused_while_open_and_backs_up_current(ui_app, tmp_path: Path) -> None:
    h = ui_app
    path = samples.standard_text(tmp_path / "Vertrag.pdf")
    before = digest(path)
    doc = open_pdf(h, path)
    doc.rotatePages([0], 90)
    settle(h)
    doc.saveDocument()
    settle(h)
    after = digest(path)
    c = refreshed(h)
    old = c.versionList[0]["path"]
    c.restoreReplace(old)  # Dokument ist offen: abgelehnt
    pump(0.5)
    assert digest(path) == after
    reader(h).closeTab(reader(h).currentKey)
    settle(h)
    c.restoreReplace(old)  # Rückfrage wird in Tests mit »Original ersetzen« beantwortet
    assert wait_until(lambda: digest(path) == before, 20)
    c = refreshed(h)
    contents = [digest(Path(row["path"])) for row in c.versionList]
    assert contents[0] == after and before in contents  # Stand vor dem Ersetzen ist gesichert
    count = len(c.versionList)
    c.deleteVersion(c.versionList[-1]["path"])
    assert wait_until(lambda: len(controller(h).versionList) == count - 1, 20)
    c.deleteDocument()
    assert wait_until(lambda: controller(h).documents == [], 20)
    assert path.is_file()


def test_auto_backup_keeps_unsaved_state_without_saving(ui_app, tmp_path: Path) -> None:
    h = ui_app
    path = samples.standard_text(tmp_path / "Entwurf.pdf")
    original = digest(path)
    doc = open_pdf(h, path)
    doc.rotatePages([0], 90)
    settle(h)
    assert doc.dirty
    c = controller(h)
    c.autoBackup()
    assert wait_until(lambda: c.lastAuto != "", 20)
    settle(h)
    assert doc.dirty and digest(path) == original  # das Dokument selbst bleibt ungespeichert
    c = refreshed(h)
    assert [row["kind"] for row in c.versionList] == ["auto"]
    import pikepdf

    with pikepdf.open(c.versionList[0]["path"]) as pdf:
        assert int(pdf.pages[0].obj.get("/Rotate", 0)) == 90
    # Ohne weitere Änderung keine neue Fassung
    c.autoBackup()
    pump(1.0)
    assert len(refreshed(h).versionList) == 1
    doc.rotatePages([0], 90)
    settle(h)
    c.autoBackup()
    assert wait_until(lambda: len(refreshed(h).versionList) == 2, 20)


def test_settings_saved_and_disabled_keeps_nothing(ui_app, tmp_path: Path) -> None:
    h = ui_app
    c = controller(h)
    h.app.openTool("versions")
    pump(0.3)
    c.setMinutes(5)
    c.setKeep(50)
    c.setDays(365)
    c.setMinutes(7)  # nicht erlaubt: bleibt
    c.setEnabled(False)
    pump(0.2)
    assert (c.minutes, c.keep, c.days, c.enabled) == (5, 50, 365, False) and not c._timer.isActive()  # noqa: SLF001
    h = neustart(h)
    c = controller(h)
    assert (c.minutes, c.keep, c.days, c.enabled) == (5, 50, 365, False)
    path = samples.standard_text(tmp_path / "Aus.pdf")
    doc = open_pdf(h, path)
    doc.rotatePages([0], 90)
    settle(h)
    doc.saveDocument()
    settle(h)
    assert refreshed(h).documents == []
    c.setEnabled(True)
    assert c._timer.isActive() and c._timer.interval() == 5 * 60_000  # noqa: SLF001
    assert not h.messages()


def _office_available() -> bool:
    from tools.sheet import engine as calc
    from tools.text import engine as text

    return calc.available() and text.available()


needs_office = pytest.mark.skipif(not _office_available(), reason="LibreOffice fehlt – Text- und Rechenkern nicht prüfbar")


@needs_office
def test_text_and_sheet_versions(ui_app, tmp_path: Path) -> None:
    """Text und Tabelle: vor dem Speichern über die Datei der bisherige Stand, automatisch eine Kopie des
    ungespeicherten Stands – das Dokument bleibt dabei »geändert« und an seinem Ort."""
    import openpyxl
    from PySide6.QtCore import QEvent, Qt
    from PySide6.QtGui import QGuiApplication, QKeyEvent

    h = ui_app
    h.app.dialogs.shutdown()
    book = openpyxl.Workbook()
    book.active.append(["Kunde", "Betrag"])
    book.save(tmp_path / "Liste.xlsx")
    path = tmp_path / "Liste.xlsx"
    first = digest(path)
    r = reader(h)
    r.open_paths([str(path)])
    assert wait_until(lambda: r.current is not None and getattr(r.current, "ready", False), 120)
    sheet = r.current
    sheet.select(1, 0, False)
    sheet.commit("Müller", 1, 0)
    assert sheet.dirty
    c = controller(h)
    c.autoBackup()
    assert wait_until(lambda: c.lastAuto != "", 60)
    assert sheet.dirty and sheet.path == str(path) and digest(path) == first
    versions_now = refreshed(h).versionList
    assert [row["kind"] for row in versions_now] == ["auto"]
    assert openpyxl.load_workbook(versions_now[0]["path"]).active["A2"].value == "Müller"
    assert sheet.save_and_wait() and wait_until(lambda: not sheet.dirty, 10)
    rows = refreshed(h).versionList
    assert rows[0]["kind"] == "vorher" and digest(Path(rows[0]["path"])) == first

    r.newText()
    assert wait_until(lambda: r.current is not None and getattr(r.current, "kind", "") == "text" and getattr(r.current, "ready", False), 120)
    doc = r.current
    pump(1.0)
    target = tmp_path / "Notiz.docx"
    doc.save(str(target))
    assert wait_until(lambda: not doc.saving and target.is_file(), 60)
    saved = digest(target)
    canvas = h.item("textCanvas")
    canvas.forceActiveFocus()
    for char in "Neu":
        for kind in (QEvent.Type.KeyPress, QEvent.Type.KeyRelease):
            QGuiApplication.sendEvent(h.window, QKeyEvent(kind, ord(char.upper()), Qt.KeyboardModifier.NoModifier, char))
        pump(0.03)
    assert wait_until(lambda: doc.dirty, 10)
    c.lastAuto = ""
    c.autoBackup()
    assert wait_until(lambda: c.lastAuto != "", 60)
    assert doc.dirty and doc.path == str(target) and digest(target) == saved
    c = refreshed(h)
    key = next(row["key"] for row in c.documents if row["name"] == "Notiz.docx")
    c.select(key)
    assert [row["kind"] for row in c.versionList] == ["auto"] and Path(c.versionList[0]["path"]).suffix == ".docx"
    assert doc.save_and_wait() and not doc.dirty
    rows = refreshed(h).versionList
    assert rows[0]["kind"] == "vorher" and digest(Path(rows[0]["path"])) == saved
