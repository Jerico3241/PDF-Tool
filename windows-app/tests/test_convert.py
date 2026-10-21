"""Umwandeln (seit 4.0 Beta 6): PDF in Text, Word und Excel – Absätze, Überschriften, Trennstriche, Tabellen, Zahlen,
Seiten ohne Text; Office in PDF (Attrappe statt Office)."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

import editorsamples as samples
from tools.pdf_editor import convert, office
from tools.pdf_editor.document import EditorDocument
from tools.pdf_editor.errors import ReadOnlyDocument


def report_pdf(path: Path, blank_page: bool = False) -> Path:
    """Bericht mit Überschrift, getrenntem Wort am Zeilenende, zwei Absätzen und einer Tabelle (+ leere Seite)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    c.setFont("Helvetica-Bold", 20)
    c.drawString(72, 780, "Jahresbericht 2026")
    c.setFont("Helvetica", 11)
    c.drawString(72, 740, "Dies ist der erste Absatz des Berichts, der über meh-")
    c.drawString(72, 726, "rere Zeilen geht und zusammengefügt werden soll.")
    c.drawString(72, 698, "Zweiter Absatz nach einer Lücke.")
    y = 658
    c.setFont("Helvetica-Bold", 11)
    for x, text in ((72, "Artikel"), (250, "Menge"), (350, "Preis")):
        c.drawString(x, y, text)
    c.setFont("Helvetica", 11)
    for name, amount, price in (("Schrauben M4", "1.200", "12,50 €"), ("Muttern", "30", "1.234,00 €"), ("Rabatt", "", "5 %")):
        y -= 16
        c.drawString(72, y, name)
        if amount:
            c.drawString(250, y, amount)
        c.drawString(350, y, price)
    c.showPage()
    if blank_page:
        c.rect(100, 100, 200, 200, fill=1)  # nur Grafik, kein Text (wie ein Scan ohne Texterkennung)
        c.showPage()
    c.save()
    return path


def test_text(tmp_path: Path) -> None:
    document = EditorDocument.open(report_pdf(tmp_path / "bericht.pdf"))
    result = convert.convert(document, tmp_path / "bericht.txt", "txt")
    text = result.path.read_text(encoding="utf-8-sig")
    assert text.startswith("Jahresbericht 2026\n\n")
    assert "über mehrere Zeilen geht und zusammengefügt werden soll.\n\nZweiter Absatz" in text  # Trennstrich aufgelöst
    assert "Artikel\tMenge\tPreis\nSchrauben M4\t1.200\t12,50 €\nMuttern\t30\t1.234,00 €" in text
    assert result.pages == 1 and result.empty == []


def test_word(tmp_path: Path) -> None:
    document = EditorDocument.open(report_pdf(tmp_path / "bericht.pdf", blank_page=True))
    result = convert.convert(document, tmp_path / "bericht.docx", "docx")
    assert result.empty == [1]
    with zipfile.ZipFile(result.path) as package:
        names = set(package.namelist())
        body = package.read("word/document.xml").decode("utf-8")
        rels = package.read("word/_rels/document.xml.rels").decode("utf-8")
    assert {"[Content_Types].xml", "_rels/.rels", "word/document.xml", "word/styles.xml", "word/media/seite1.jpg"} <= names
    assert '<w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t xml:space="preserve">Jahresbericht 2026' in body
    assert "über mehrere Zeilen" in body
    assert body.count("<w:tr>") == 4 and body.count("<w:tc>") == 12  # Tabelle 4 × 3, auch die leere Zelle
    assert '<w:br w:type="page"/>' in body and 'r:embed="rIdImage1"' in body and 'Id="rIdImage1"' in rels  # leere Seite als Bild
    import xml.dom.minidom

    xml.dom.minidom.parseString(body)  # wohlgeformt


def test_excel(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    document = EditorDocument.open(report_pdf(tmp_path / "bericht.pdf"))
    result = convert.convert(document, tmp_path / "bericht.xlsx", "xlsx")
    sheet = load_workbook(result.path).worksheets[0]
    assert sheet.title == "Seite 1"
    rows = [row for row in sheet.iter_rows(values_only=True)]
    header = rows.index(("Artikel", "Menge", "Preis"))
    assert rows[header + 1] == ("Schrauben M4", 1200, 12.5)
    assert rows[header + 2] == ("Muttern", 30, 1234)
    assert rows[header + 3] == ("Rabatt", None, 0.05)
    assert sheet.cell(header + 2, 3).number_format == '#,##0.00 "€"' and sheet.cell(header + 4, 3).number_format == "0%"
    assert sheet.cell(header + 1, 1).font.bold


@pytest.mark.parametrize(("text", "value"), [("1.234,56", 1234.56), ("-12", -12), ("12,5 %", 0.125), ("1,234.50", 1234.5), ("3.5", 3.5), ("2026-10", None), ("Seite 3", None), ("1.2.3", None),
                                            ("0,5", 0.5), ("01067", None), ("007", None), ("+49 221 123", None)])
def test_numbers(text: str, value) -> None:
    parsed, _fmt = convert.number(text)
    assert parsed == pytest.approx(value) if value is not None else parsed is None


@pytest.mark.parametrize(("text", "fmt"), [("2026", "General"), ("1.234", "#,##0"), ("1234,56", "0.00"), ("1.234,56", "#,##0.00"), ("12 €", '#,##0 "€"')])
def test_thousands_separator_only_when_the_number_had_one(text: str, fmt: str) -> None:
    """Eine Jahreszahl »2026« erschien in Excel als »2.026« (bis 4.0 Beta 7)."""
    assert convert.number(text)[1] == fmt


def test_excel_text_starting_with_equals_is_no_formula(tmp_path: Path) -> None:
    from openpyxl import load_workbook
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path = tmp_path / "formel.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setFont("Helvetica", 11)
    c.drawString(72, 780, "=== Summe ===")
    c.drawString(72, 760, "=A1+1")
    c.save()
    result = convert.convert(EditorDocument.open(path), tmp_path / "formel.xlsx", "xlsx")
    sheet = load_workbook(result.path).worksheets[0]
    cells = [cell for row in sheet.iter_rows() for cell in row if cell.value]
    assert [cell.value for cell in cells] == ["=== Summe ===", "=A1+1"]
    assert all(cell.data_type == "s" for cell in cells)


def test_copy_protected_is_refused(tmp_path: Path) -> None:
    document = EditorDocument.open(samples.copy_protected(tmp_path / "geschuetzt.pdf"))
    with pytest.raises(ReadOnlyDocument):
        convert.convert(document, tmp_path / "x.docx", "docx")
    assert not (tmp_path / "x.docx").exists()


def test_office_to_pdf_with_fake_office(tmp_path: Path, monkeypatch) -> None:
    calls = []

    def fake(kind, source, target):
        calls.append((kind, source.name))
        samples.standard_text(target)

    monkeypatch.setattr(office, "TEST_BACKEND", fake)
    monkeypatch.setattr(office, "TEST_INSTALLED", {"word"})
    source = tmp_path / "Brief.docx"
    source.write_bytes(b"PK")
    assert office.to_pdf(source, tmp_path / "Brief.pdf").is_file()
    assert calls == [("word", "Brief.docx")]
    sheet = tmp_path / "Liste.xlsx"
    sheet.write_bytes(b"PK")
    with pytest.raises(office.OfficeError, match="Excel"):  # Excel nicht installiert
        office.to_pdf(sheet, tmp_path / "Liste.pdf")
    with pytest.raises(office.OfficeError, match="bereits"):
        office.to_pdf(source, tmp_path / "Brief.pdf")
    assert office.kind_of("a.PPTX") == "powerpoint" and office.kind_of("a.pdf") is None
