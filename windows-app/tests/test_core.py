"""Tests ohne Oberfläche: PDF-Engine, Einstellungen, Farben (Design-Tokens), Mica-Rezept."""

from __future__ import annotations

import json
from pathlib import Path

import pytest

import appstate
import design as theme
import engine
import mica


# --- Engine ------------------------------------------------------------------------


def test_pruefe_excel_counts_active_contracts(excel_file: Path) -> None:
    result = engine.pruefe_excel(excel_file, [{"enthaelt": "Hott-KI", "zyklus": "jährlich"}])
    assert result["ok"] is True
    assert result["text"].startswith("3 aktive Verträge")
    assert "1 inaktiv ausgeblendet" in result["text"]
    assert result["kunden"] == ["10042"]
    assert result["firmen"] == ["Muster GmbH"]
    assert result["mails"] == ["rechnung@muster.de"]
    zyklus = {zeile["nr"]: zeile["zyklus"] for zeile in result["zeilen"]}
    assert zyklus["V-1002"] == "jährlich"
    assert zyklus["V-1001"] == "jeden Monat"


def test_pruefe_excel_missing_file(tmp_path: Path) -> None:
    result = engine.pruefe_excel(tmp_path / "fehlt.xlsx")
    assert result["ok"] is False
    assert "nicht gefunden" in result["text"]


def test_erstelle_pdf_with_header_footer(excel_file: Path, tmp_path: Path) -> None:
    pypdf = pytest.importorskip("pypdf")
    auftrag = engine.PdfAuftrag(
        excel=excel_file,
        logo=appstate.DEFAULT_LOGO,
        kundennummer="10042",
        firmenname="Muster GmbH",
        zielordner=tmp_path,
        dateiname="Uebersicht_{kunde}_Kd{kd}.pdf",
        seitenformat="quer",
        kopfzeile="Kopf {kd}",
        fusszeile="Fuß {kunde}",
        regeln=[{"enthaelt": "Hott-KI", "zyklus": "jährlich"}],
    )
    pfad = engine.erstelle_pdf(auftrag)
    assert pfad == tmp_path / "Uebersicht_Muster_GmbH_Kd10042.pdf"
    reader = pypdf.PdfReader(str(pfad))
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "Kopf 10042" in text
    # Im Text erscheint der Firmenname wie eingetragen; nur der Dateiname ist dateisystemtauglich.
    assert "Fuß Muster GmbH" in text
    assert "Supportvertrag" in text
    assert "V-0999" not in text  # inaktiv
    assert "Seite 1 von" in text
    width, height = reader.pages[0].mediabox.width, reader.pages[0].mediabox.height
    assert width > height  # Querformat


def test_erstelle_pdf_requires_customer_number(excel_file: Path, tmp_path: Path) -> None:
    with pytest.raises(ValueError):
        engine.erstelle_pdf(engine.PdfAuftrag(excel=excel_file, logo=appstate.DEFAULT_LOGO, kundennummer="", zielordner=tmp_path))


# --- Standard-Fußzeile -----------------------------------------------------------------------

ERWARTETE_FUSSZEILE = (
    "Die oben aufgeführte Auflistung gibt den aktuellen Stand Ihrer Verträge sowie die derzeit geltenden Vertragspreise wieder.\n"
    "Alle genannten Preise verstehen sich zuzüglich der jeweils geltenden gesetzlichen Mehrwertsteuer.\n"
    "\n"
    "Die Angaben erfolgen gemäß Ihren abgeschlossenen Verträgen sowie den jeweils geltenden Vertragsbedingungen und berücksichtigen gegebenenfalls bereits erfolgte Preisanpassungen."
)


def test_default_footer_is_exact() -> None:
    assert appstate.DEFAULT_FOOTER == ERWARTETE_FUSSZEILE


def test_default_footer_has_single_source() -> None:
    root = Path(__file__).resolve().parents[1] / "app"
    fundstellen = [p.relative_to(root).as_posix() for p in root.rglob("*.py") if "Mehrwertsteuer" in p.read_text(encoding="utf-8")]
    assert fundstellen == ["appstate.py"]


@pytest.mark.parametrize(
    "eintrag,erwartet",
    [
        ({}, ERWARTETE_FUSSZEILE),  # Schlüssel fehlt (erster Start, alte Vorlage)
        ({"fusszeile": None}, ERWARTETE_FUSSZEILE),  # null
        ({"fusszeile": ""}, ERWARTETE_FUSSZEILE),  # leerer Standardwert älterer Versionen
        ({"fusszeile": "  \n "}, ERWARTETE_FUSSZEILE),
        ({"fusszeile": 42}, ERWARTETE_FUSSZEILE),  # ungültiger Wert
        ({"fusszeile": "Eigener Text\nZeile 2\n\nZeile 4"}, "Eigener Text\nZeile 2\n\nZeile 4"),  # alte eigene Fußzeile bleibt
        ({"fusszeile": "  Einzug\n\n", "fusszeile_explizit": True}, "  Einzug\n\n"),  # exakt, keine Kürzung
        ({"fusszeile": "", "fusszeile_explizit": True}, ""),  # bewusst leer gespeichert
    ],
)
def test_footer_migration_rules(eintrag: dict, erwartet: str) -> None:
    assert appstate.footer_from(eintrag) == erwartet


def test_customer_footer_keeps_current_for_empty_values() -> None:
    from tools.contract_overview.customers.texts import footer_of
    from tools.contract_overview.customers.migration import customer_from_legacy

    for leer in ({"fusszeile": ""}, {"fusszeile": None}, {}, {"fusszeile": "  \n "}):
        kunde = customer_from_legacy({"firmenname": "A", **leer}, "")
        assert kunde.footer is None and footer_of(kunde) is None
    kunde = customer_from_legacy({"firmenname": "A", "fusszeile": "Kunde A\n\nGruß"}, "")
    assert footer_of(kunde).text == "Kunde A\n\nGruß"


def _textzeilen(page) -> list[tuple[float, str]]:
    """Textzeilen einer PDF-Seite mit absoluter y-Position (pt, von unten)."""
    zeilen: list[tuple[float, str]] = []

    def visit(text, cm, tm, _font, _size):
        if text.strip():
            zeilen.append((round(tm[5] * cm[3] + cm[5], 1), " ".join(text.split())))

    page.extract_text(visitor_text=visit)
    return zeilen


@pytest.mark.parametrize("seitenformat", ["hoch", "quer"])
def test_pdf_contains_default_footer_with_paragraph(excel_file: Path, tmp_path: Path, seitenformat: str) -> None:
    pypdf = pytest.importorskip("pypdf")
    pfad = engine.erstelle_pdf(
        engine.PdfAuftrag(excel=excel_file, logo=appstate.DEFAULT_LOGO, kundennummer="10042", zielordner=tmp_path, seitenformat=seitenformat, fusszeile=appstate.DEFAULT_FOOTER)
    )
    seite = pypdf.PdfReader(str(pfad)).pages[0]
    text = " ".join(seite.extract_text().split())
    for absatz in ERWARTETE_FUSSZEILE.split("\n"):
        if absatz:
            assert absatz in text
    zeilen = _textzeilen(seite)

    def y_von(anfang: str) -> float:
        return next(y for y, t in zeilen if t.startswith(anfang))

    y1 = y_von("Die oben aufgeführte Auflistung")
    y2 = y_von("Alle genannten Preise")
    y3 = y_von("Die Angaben erfolgen")
    # Zeilenabstand 10 pt, zwischen zweitem und drittem Block eine Leerzeile (20 pt)
    assert y1 - y2 == pytest.approx(10, abs=0.6)
    assert y2 - y3 == pytest.approx(20, abs=0.6)
    # Die Tabelle endet oberhalb der Fußzeile (keine Überdeckung)
    tabelle_unten = min(y for y, t in zeilen if t in ("V-1001", "V-1002", "V-1003"))
    assert tabelle_unten > y1 + 8


def test_pdf_footer_keeps_user_formatting(excel_file: Path, tmp_path: Path) -> None:
    pypdf = pytest.importorskip("pypdf")
    eigene = "Eigener kundenspezifischer Text\nZeile 2\n\nZeile 4\n"
    pfad = engine.erstelle_pdf(engine.PdfAuftrag(excel=excel_file, logo=appstate.DEFAULT_LOGO, kundennummer="10042", zielordner=tmp_path, fusszeile=eigene))
    zeilen = _textzeilen(pypdf.PdfReader(str(pfad)).pages[0])
    y = {t: y for y, t in zeilen if t in ("Eigener kundenspezifischer Text", "Zeile 2", "Zeile 4")}
    assert y["Eigener kundenspezifischer Text"] - y["Zeile 2"] == pytest.approx(10, abs=0.6)
    assert y["Zeile 2"] - y["Zeile 4"] == pytest.approx(20, abs=0.6)


# --- Einstellungen (kompatibel zu 2.0.5) ---------------------------------------------------


def test_config_roundtrip_is_atomic(tmp_path: Path) -> None:
    target = tmp_path / "gui-config.json"
    assert appstate.save_config({"a": 1, "ä": "ö"}, target)
    assert appstate.load_config(target) == {"a": 1, "ä": "ö"}
    assert not [p for p in tmp_path.iterdir() if p.name.endswith(".tmp")]


def test_config_corrupt_returns_empty(tmp_path: Path) -> None:
    target = tmp_path / "gui-config.json"
    target.write_text("{kaputt", encoding="utf-8")
    assert appstate.load_config(target) == {}


def test_state_reads_205_config() -> None:
    cfg = {
        "firmenname": "Alt GmbH",
        "kunden": [{"firmenname": "A", "kundennummer": "1"}, "müll"],
        "bausteine": [{"name": "B1", "text": "T"}, {"text": "ohne Name"}],
        "regeln": [{"enthaelt": "X", "zyklus": "Y"}, {"enthaelt": ""}],
        "pdfs": ["/nicht/da.pdf", 3],
        "vorlagen": [{"name": "V1"}],
        "staende": {"k": 1},
        "gesehen": "2.0.5",
        "theme": "dark",
        "accent": "#c42b1e",
    }
    state = appstate.State(cfg)
    assert not hasattr(state, "kunden")  # übernimmt die Kundenakte 2.0 (siehe test_customers.py)
    assert [b["name"] for b in state.bausteine] == ["B1"]
    assert state.regeln == [{"enthaelt": "X", "zyklus": "Y"}]
    assert state.pdfs == ["/nicht/da.pdf"]
    assert state.staende == {"k": 1}
    assert appstate.major_minor(state.gesehen) != appstate.major_minor(appstate.VERSION)


def test_state_default_rules_when_missing() -> None:
    assert appstate.State({}).regeln == [{"enthaelt": "Hott-KI", "zyklus": "jährlich"}]
    assert appstate.State({"regeln": []}).regeln == []


def test_legacy_history_entries_become_customers() -> None:
    from tools.contract_overview.customers.migration import customers_from_legacy

    eintraege = [{"firmenname": f"Firma {n}", "kundennummer": str(n), "rechnungsempfaenger": f"r{n}@x.de"} for n in range(12)]
    kunden, skipped = customers_from_legacy(eintraege + [{"firmenname": "", "kundennummer": ""}, "müll"])
    assert len(kunden) == 12 and skipped == 2
    assert len({k.id for k in kunden}) == 12 and kunden[3].emails == ["r3@x.de"]


def test_templates_blocks_rules() -> None:
    state = appstate.State({})
    state.save_vorlage({"name": "A", "titel": "1"})
    state.save_vorlage({"name": "B"})
    state.save_vorlage({"name": "A", "titel": "2"})
    assert [v["name"] for v in state.vorlagen] == ["A", "B"]
    assert state.find_vorlage("A")["titel"] == "2"
    assert state.delete_vorlage("A") and not state.delete_vorlage("A")
    state.save_baustein("F", "Text")
    assert state.find_baustein("F")["text"] == "Text"
    assert state.delete_baustein("F")
    state.add_regel("hott-ki", "halbjährlich")
    assert state.regeln == [{"enthaelt": "hott-ki", "zyklus": "halbjährlich"}]
    assert state.delete_regel(0)["enthaelt"] == "hott-ki"
    assert state.delete_regel(5) is None


def test_versions_are_consistent() -> None:
    root = Path(__file__).resolve().parents[1]
    version = (root / "VERSION").read_text(encoding="utf-8").strip()
    iss = (root / "installer" / "PDF-Tool.iss").read_text(encoding="utf-8-sig")
    readme = (root / "README.txt").read_text(encoding="utf-8")
    assert appstate.VERSION == version == "4.0.0"
    # Das Setup liest die Version aus derselben Datei, statt sie zu wiederholen.
    assert r'FileOpen(AddBackslash(SourcePath) + "..\VERSION")' in iss
    assert "2.0.5" not in iss.split("[Setup]")[1].split("[Languages]")[0]
    assert f"Version {version}" in readme
    assert f"PDF-Tool-Setup-{version}.exe" in readme
    assert not (root / "installer" / "main.go").exists()
    assert not (root / "installer" / "go.mod").exists()
    # Die aktuelle Version steht nirgends hart codiert – weder im App-Code noch auf der Downloadseite
    # (Hinweise auf frühere Versionen wie »bis 2.0.5« sind Geschichte, keine Versionsanzeige).
    for datei in list((root / "app").rglob("*.py")) + list((root / "app" / "qml").rglob("*.qml")):
        if "__pycache__" not in datei.parts:
            assert version not in datei.read_text(encoding="utf-8"), datei.name
    seite = (root.parent / "src" / "components" / "landing-page.tsx").read_text(encoding="utf-8")
    assert "windows-app/VERSION?raw" in seite and version not in seite
    assert (root / "release-notes" / f"{version}.md").is_file()


def test_readme_stays_a_short_introduction() -> None:
    """README-Regel: Die README.md stellt PDF Tool kurz vor – keine Versionsnummer, die bei jedem
    Release geändert werden müsste, keine Release-Historie, keine Test- oder CI-Berichte."""
    import re

    root = Path(__file__).resolve().parents[1]
    version = (root / "VERSION").read_text(encoding="utf-8").strip()
    haupt = (root.parent / "README.md").read_text(encoding="utf-8")
    assert len(haupt.splitlines()) <= 80
    assert version not in haupt
    assert not re.search(r"(?i)neu in (version )?\d", haupt)
    assert [line for line in haupt.splitlines() if line.startswith("## ")] == ["## Funktionen", "## Oberfläche", "## Datenschutz", "## Installation", "## Entwicklung", "## Lizenz"]
    for link in re.findall(r"\]\(((?:docs|windows-app)/[^)#]+)\)", haupt) + ["THIRD_PARTY_LICENSES.md"]:
        assert (root.parent / link).exists(), link


# SHA-256 des Hottgenroth-Logos der Vertragsübersichten (Dokument-Branding, kein App-Branding)
HOTT_LOGO_SHA256 = "a1564c62fe01caa4cb041d8178578da051449d87db187b3cd7de2b667d202c5d"


def test_branding_is_consistent() -> None:
    import hashlib

    import winsys as windows

    root = Path(__file__).resolve().parents[1]
    iss = (root / "installer" / "PDF-Tool.iss").read_text(encoding="utf-8-sig")
    assert not (root / "installer" / "Uebersichten-Ersteller.iss").exists()
    assert '#define AppName        "PDF Tool"' in iss
    # Die AppId bleibt – ein Update vom Übersichten-Ersteller erzeugt keinen zweiten App-Eintrag.
    assert '#define AppGuid        "59C40061-D5E5-446D-ACB4-E077D3C71E1A"' in iss
    assert "OutputBaseFilename=PDF-Tool-Setup-{#AppVersion}" in iss
    assert f'#define AppUserModelID "{windows.APP_USER_MODEL_ID}"' in iss
    assert f'#define AppMutexName   "{windows.APP_MUTEX}"' in iss
    assert "AppMutex={#AppMutexName},{#OldMutexName}" in iss  # läuft noch die alte App?
    assert appstate.APP_NAME == "PDF Tool" and appstate.DATA_FOLDER == "PDF-Tool"
    assert appstate.LEGACY_DATA_FOLDER == "Uebersichten-Ersteller"
    # Alte Produktnamen stehen nur noch dort, wo die Vorversion übernommen wird.
    for datei in list((root / "app").rglob("*.py")) + list((root / "app" / "qml").rglob("*.qml")):
        if "__pycache__" in datei.parts:
            continue
        text = datei.read_text(encoding="utf-8")
        for alt in ("PDF Multi Tool", "Vertragsübersichten-Ersteller", "VertraView"):
            assert alt not in text, datei
        if datei.name != "appstate.py":
            assert "Übersichten-Ersteller" not in text and "Uebersichten-Ersteller" not in text, datei
    # Das Logo in den erzeugten Vertragsübersichten ist Dokument-Branding und bleibt unverändert.
    assert appstate.DEFAULT_LOGO.name == "hott_logo_final.png"
    assert hashlib.sha256(appstate.DEFAULT_LOGO.read_bytes()).hexdigest() == HOTT_LOGO_SHA256


# --- Farben und Bewegung -------------------------------------------------------------------


@pytest.mark.parametrize("name,value", theme.ACCENTS)
def test_accent_contrast(name: str, value: str) -> None:
    accent = theme.derive_accent(value)
    light = theme.build_palette(False, accent)
    dark = theme.build_palette(True, accent)
    assert theme.contrast(light.accent, light.on_accent) >= 4.5, name
    assert theme.contrast(dark.accent, dark.on_accent) >= 4.5, name


def test_palettes_text_contrast() -> None:
    accent = theme.derive_accent(theme.DEFAULT_ACCENT)
    for dark in (False, True):
        pal = theme.build_palette(dark, accent)
        assert theme.contrast(pal.text, pal.card) >= 7
        assert theme.contrast(pal.text2, pal.card) >= 4.5
        assert pal.card != pal.layer != pal.mica


def test_mica_tint_is_subtle() -> None:
    for pixel in ((0, 90, 200), (220, 30, 30), (255, 255, 255), (0, 0, 0)):
        light = mica.tint_pixel(pixel, dark=False)
        dark = mica.tint_pixel(pixel, dark=True)
        assert all(abs(c - 243) <= 14 for c in light), (pixel, light)
        assert all(abs(c - 32) <= 18 for c in dark), (pixel, dark)


def test_config_file_json_is_valid(tmp_path: Path) -> None:
    target = tmp_path / "c.json"
    appstate.save_config({"liste": [1, 2]}, target)
    assert json.loads(target.read_text(encoding="utf-8")) == {"liste": [1, 2]}


def test_german_date_text_is_read_day_first():
    engine._lade_pandas()
    import datetime

    import pandas as pd

    werte = pd.Series(["01.02.2021", "05.03.2022", "13.02.2021", datetime.datetime(2020, 4, 5), None, "", "2021-07-08"])
    gelesen = [None if pd.isna(stamp) else stamp.strftime("%Y-%m-%d") for stamp in engine.datum_spalte(werte)]
    assert gelesen == ["2021-02-01", "2022-03-05", "2021-02-13", "2020-04-05", None, None, "2021-07-08"]
    assert engine.fmt_date("01.02.2021") == "01.02.2021"


def test_merged_cells_apply_to_every_contract_row(tmp_path: Path) -> None:
    # Anwenderstatus, Kundennummer, Firma und E-Mail einmal verbunden für drei Verträge: alle drei sind aktiv
    from datetime import datetime

    from openpyxl import load_workbook

    from conftest import write_excel

    rows = [[f"V-{n}", datetime(2021, n, 1), "monatlich", 10.0, "Lastschr", f"Pflege {n}", None, None, None, None] for n in (1, 2, 3)]
    rows[0][6:] = ["rechnung@muster.de", 10042, "Muster GmbH", "Aktiv"]
    path = write_excel(tmp_path / "verbunden.xlsx", rows)
    book = load_workbook(path)
    for column in "GHIJ":
        book.active.merge_cells(f"{column}2:{column}4")
    book.save(path)
    result = engine.pruefe_excel(path)
    assert result["text"].startswith("3 aktive Verträge") and "ausgeblendet" not in result["text"]
    assert result["kunden"] == ["10042"] and result["firmen"] == ["Muster GmbH"] and result["mails"] == ["rechnung@muster.de"]


def test_unformatted_excel_serial_date_is_a_date() -> None:
    # »Beginnt am« ohne Datumsformat: Excel liefert 44197 statt 01.01.2021 – nicht 01.01.1970
    engine._lade_pandas()
    assert engine.fmt_date(engine.als_datum(44197)) == "01.01.2021"
    assert engine.fmt_date(engine.als_datum(45292.0)) == "01.01.2024"
    assert engine.fmt_date(engine.als_datum(float("nan"))) == "–"


@pytest.mark.parametrize(
    ("muster", "erwartet"),
    [
        ("Vertrag_{Kd}.pdf", "Vertrag_10042.pdf"),
        ("Vertrag {2026}.pdf", "Vertrag {2026}.pdf"),
        ("Vertrag_{kd.pdf", "Vertrag_{kd.pdf"),
        ("Vertrag_{kd}_{unbekannt}", "Vertrag_10042_{unbekannt}.pdf"),
        ("Uebersicht_{kunde}_Kd{kd}.pdf", "Uebersicht_Muster_GmbH_Kd10042.pdf"),
    ],
)
def test_file_name_pattern_replaces_only_known_placeholders(muster: str, erwartet: str) -> None:
    # unbekannte oder kaputte Klammern bleiben als Text stehen, statt jede PDF scheitern zu lassen
    assert engine.dateiname_fuer(muster, "10042", "Muster GmbH") == erwartet
