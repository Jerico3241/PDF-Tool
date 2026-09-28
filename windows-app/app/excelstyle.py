"""Zellformate aus Excel-Dateien – in Version 2.2 ausschließlich Fettschrift.

pandas bleibt für Werte, Spalten, Datumswerte, Filter und Sortierung zuständig.
Diese Schicht liest zusätzlich die Formatierung derselben Zellen:

1. Arbeitsmappe einmal öffnen (openpyxl für .xlsx/.xlsm, xlrd mit
   ``formatting_info`` für .xls)
2. Formatierung der benötigten Zellen auslesen
3. Zuordnung (Excel-Zeile, Excel-Spalte) → ``ExcelCellStyle`` erstellen
4. Arbeitsmappe schließen

Jede Datenzeile trägt ihre ursprüngliche Excel-Zeilennummer in der Spalte
``_source_excel_row``; Filtern, Sortieren und ``reset_index`` ändern daran nichts.
Bevor Stile verwendet werden, wird die Zuordnung an den Vertragsnummern geprüft.
Passt sie nicht oder liefert eine Datei keine Stile (exotische XLS), bleibt
alles normal formatiert – die PDF entsteht trotzdem.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Mapping

SOURCE_ROW = "_source_excel_row"

_XLSX_MAGIC = b"PK\x03\x04"
_XLS_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


@dataclass(frozen=True)
class ExcelCellStyle:
    """Formatierung einer Excel-Zelle, soweit sie in die PDF übernommen wird.

    Weitere Merkmale (z. B. kursiv) lassen sich später hier ergänzen.
    """

    bold: bool = False


NORMAL = ExcelCellStyle()
BOLD = ExcelCellStyle(bold=True)


class ExcelStyles:
    """Zuordnung (Excel-Zeile, Excel-Spalte) → ``ExcelCellStyle``; beide 1-basiert wie in Excel.

    Gespeichert werden nur Zellen, die vom normalen Stil abweichen.
    """

    def __init__(self, cells: Mapping[tuple[int, int], ExcelCellStyle] | None = None, available: bool = False, reason: str = "") -> None:
        self._cells = {key: style for key, style in (cells or {}).items() if style != NORMAL}
        self.available = available
        self.reason = reason

    def get(self, row, column) -> ExcelCellStyle:
        try:
            key = (int(row), int(column))
        except (TypeError, ValueError):
            return NORMAL
        return self._cells.get(key, NORMAL)

    def bold(self, row, column) -> bool:
        return self.get(row, column).bold

    def bold_cells(self) -> int:
        return sum(1 for style in self._cells.values() if style.bold)

    def __len__(self) -> int:
        return len(self._cells)


def unavailable(reason: str) -> ExcelStyles:
    return ExcelStyles(available=False, reason=reason)


def cell_key(value) -> str:
    """Vergleichbarer Text eines Zellwerts (pandas und Excel-Bibliotheken liefern Zahlen unterschiedlich)."""
    if value is None:
        return ""
    if isinstance(value, float):
        if value != value:  # NaN
            return ""
        if value.is_integer():
            return str(int(value))
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].lstrip("-").isdigit():
        return text[:-2]
    return text


def file_kind(path: Path) -> str:
    """``xlsx`` oder ``xls`` anhand des Dateiinhalts (wie pandas), sonst ``""``."""
    try:
        with open(path, "rb") as handle:
            head = handle.read(8)
    except OSError:
        return ""
    if head.startswith(_XLSX_MAGIC):
        return "xlsx"
    if head.startswith(_XLS_MAGIC):
        return "xls"
    return ""


def read_styles(
    path: Path,
    rows: Iterable[int],
    columns: Iterable[int],
    expected: Mapping[tuple[int, int], str] | None = None,
    sheet_index: int = 0,
) -> ExcelStyles:
    """Liest die Formatierung der Zellen ``rows`` × ``columns`` (Excel-Nummern, 1-basiert).

    ``expected`` ordnet einzelnen Zellen ihren erwarteten Text zu (z. B. die
    Vertragsnummer jeder Zeile). Weicht eine davon ab, stimmt die Zuordnung
    nicht – dann werden keine Stile verwendet.
    """
    wanted_rows = sorted({int(r) for r in rows if r is not None})
    wanted_cols = sorted({int(c) for c in columns if c is not None})
    if not wanted_rows or not wanted_cols:
        return ExcelStyles(available=True)
    kind = file_kind(Path(path))
    try:
        if kind == "xlsx":
            cells, values = _read_openpyxl(Path(path), wanted_rows, wanted_cols, sheet_index)
        elif kind == "xls":
            cells, values = _read_xlrd(Path(path), wanted_rows, wanted_cols, sheet_index)
        else:
            return unavailable("Unbekanntes Dateiformat")
    except Exception as exc:  # noqa: BLE001 - Formatierung ist optional, die PDF entsteht trotzdem
        return unavailable(f"Formatierung nicht lesbar: {exc.__class__.__name__}")
    for key, text in (expected or {}).items():
        if values.get(key, "") != text:
            return unavailable("Zuordnung der Zeilen nicht eindeutig")
    return ExcelStyles(cells, available=True)


def _read_openpyxl(path: Path, rows: list[int], cols: list[int], sheet_index: int):
    from openpyxl import load_workbook

    # Wie pandas: nur lesen, berechnete Werte statt Formeln
    book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = book.worksheets[sheet_index]
        wanted_rows = set(rows)
        wanted_cols = set(cols)
        cells: dict[tuple[int, int], ExcelCellStyle] = {}
        values: dict[tuple[int, int], str] = {}
        min_col, max_col = cols[0], cols[-1]
        for offset, row in enumerate(sheet.iter_rows(min_row=rows[0], max_row=rows[-1], min_col=min_col, max_col=max_col)):
            row_number = rows[0] + offset
            if row_number not in wanted_rows:
                continue
            for col_offset, cell in enumerate(row):
                column = min_col + col_offset
                if column not in wanted_cols:
                    continue
                font = getattr(cell, "font", None)
                if font is not None and bool(getattr(font, "b", False)):
                    cells[(row_number, column)] = BOLD
                values[(row_number, column)] = cell_key(getattr(cell, "value", None))
        return cells, values
    finally:
        book.close()


def _read_xlrd(path: Path, rows: list[int], cols: list[int], sheet_index: int):
    import xlrd

    book = xlrd.open_workbook(str(path), formatting_info=True, on_demand=True)
    try:
        sheet = book.sheet_by_index(sheet_index)
        cells: dict[tuple[int, int], ExcelCellStyle] = {}
        values: dict[tuple[int, int], str] = {}
        for row_number in rows:
            rowx = row_number - 1
            if rowx >= sheet.nrows:
                continue
            for column in cols:
                colx = column - 1
                if colx >= sheet.row_len(rowx):
                    continue
                xf = book.xf_list[sheet.cell_xf_index(rowx, colx)]
                font = book.font_list[xf.font_index]
                if bool(getattr(font, "bold", False)) or int(getattr(font, "weight", 400) or 400) >= 700:
                    cells[(row_number, column)] = BOLD
                values[(row_number, column)] = cell_key(sheet.cell_value(rowx, colx))
        return cells, values
    finally:
        book.release_resources()
